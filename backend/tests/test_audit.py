"""End-to-end audit against a local fixture site with known defects."""
import io

import openpyxl

from app.worker import run_once


def _audit(org, site_server, **body):
    p = org.project("Acme Plumbing")
    site = org.post("/sites", {"project_id": p["id"], "domain": site_server}).json()
    r = org.post(f"/sites/{site['id']}/audits", {"max_pages": 50, **body})
    assert r.status_code == 200, r.text
    run_once()
    return p, site, org.get(f"/audits/{r.json()['id']}").json()


def test_full_audit_pipeline(org, site_server):
    p, site, a = _audit(org, site_server)
    assert a["status"] == "COMPLETED", a
    assert len(a["results"]) == 62
    g = {r["code"]: r for r in a["results"]}
    expect = {
        "CHK-05": "pass",      # robots.txt with Sitemap line
        "CHK-04": "fail",      # /private/ in sitemap but disallowed
        "CHK-03": "fail",      # noindex water-heaters page is in the sitemap
        "CHK-12": "fail",      # link to /missing.html
        "CHK-13": "fail",      # orphan.html only in sitemap
        "CHK-16": "fail",      # duplicate titles (blog posts)
        "CHK-21": "fail",      # two H1s on drains page
        "CHK-26": "fail",      # identical blog posts
        "CHK-34": "fail",      # /old-page 302 → 301 chain
        "CHK-35": "fail",      # 302 hop
        "CHK-37": "fail",      # 404 crawled
        "CHK-25": "pass", "CHK-60": "pass", "CHK-61": "pass", "CHK-14": "pass", "CHK-08": "pass",
        "CHK-09": "not_run", "CHK-43": "not_run", "CHK-57": "not_run", "CHK-62": "not_run",
        "CHK-27": "manual", "CHK-53": "manual", "CHK-52": "na",
    }
    wrong = {k: (g[k]["grade"], v, g[k]["message"]) for k, v in expect.items() if g[k]["grade"] != v}
    assert not wrong, wrong
    assert 0 < a["score"] < 100
    assert a["summary"]["issues"]["critical"] >= 1 and a["summary"]["pages_crawled"] >= 8
    assert set(a["category_scores"]) >= {"Indexation", "On-Page", "Technical"}
    pri = [i["priority_score"] for i in a["issues"]]
    assert pri == sorted(pri, reverse=True)
    # orphan URL is evidenced
    orphan = next(i for i in a["issues"] if i["rule"] == "CHK-13")
    assert any(u.endswith("/orphan.html") for u in orphan["affected_urls"])
    # NOT CONNECTED messaging is explicit, not faked
    assert "NOT CONNECTED" in g["CHK-59"]["message"]

    # action plan → tasks
    plan = org.get(f"/action-plans/{a['plan_id']}").json()
    assert len(plan["items"]) == len(a["issues"])
    assert org.post(f"/action-plans/{a['plan_id']}/tasks", {}).json()["created"] == len(a["issues"])
    assert org.post(f"/action-plans/{a['plan_id']}/tasks", {}).status_code == 409
    tasks = org.get(f"/tasks?project_id={p['id']}").json()
    assert any(t["title"].startswith("[CHK-") for t in tasks)

    # reports
    for fmt in ("pdf", "xlsx"):
        rid = org.post("/reports", {"audit_run_id": a["id"], "format": fmt}).json()["id"]
        run_once()
        r = org.get(f"/reports/{rid}/download")
        assert r.status_code == 200
        if fmt == "pdf":
            assert r.content[:4] == b"%PDF"
        else:
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
            assert wb.sheetnames == ["Summary", "Issues", "Checklist", "Affected URLs"]
            assert wb["Checklist"].max_row == 63
    usage = {x["key"]: x for x in org.get("/billing/usage").json()["meters"]}
    assert usage["audit_pages"]["used"] == a["summary"]["pages_crawled"]  # unused reserve refunded


def test_crawler_refuses_private_targets_by_default(monkeypatch):
    from app.services import crawler
    monkeypatch.setattr(crawler.S, "allow_private_crawl", False)
    for u in ("http://127.0.0.1/", "http://169.254.169.254/latest/meta-data/", "http://localhost:5432/", "http://100.100.100.200/"):
        try:
            crawler.assert_public(u)
            raise AssertionError(f"{u} should be blocked")
        except crawler.CrawlBlocked:
            pass


def test_competitor_audit_and_compare(org, site_server):
    p, site, a = _audit(org, site_server)
    c = org.post("/competitors", {"project_id": p["id"], "domain": site_server}).json()
    r = org.post(f"/sites/{site['id']}/audits", {"max_pages": 20, "competitor_id": c["id"]}).json()
    run_once()
    cmp = org.get(f"/audits/{a['id']}/compare/{r['id']}").json()
    assert cmp["b"]["domain"] == site_server and len(cmp["checks"]) == 62
