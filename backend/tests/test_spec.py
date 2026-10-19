"""The workbook is the source of truth: every sheet is read and every row reaches the product."""
import copy
import io

import openpyxl

from app.services.checks import REG
from app.spec import traceability
from app.spec.parser import DEFAULT_WORKBOOK, REQUIRED_SHEETS, diff_specs, parse_workbook, public_spec


def test_all_nine_sheets_parsed():
    s = parse_workbook()
    assert s["meta"]["sheets_read"] == REQUIRED_SHEETS
    c = s["meta"]["counts"]
    assert c == {"modules": 76, "features": 62, "tables": 51, "integrations": 23, "phases": 7, "checks": 62, "pricing_rows": 26}
    assert s["pricing"]["tiers"] == ["Essentials", "Advanced", "Enterprise"]
    assert s["roadmap"]["totals"]["Sequential total (weeks)"] == 58
    assert s["cost_estimate"]["summary"]["grand_total"] == 435160


def test_cost_estimate_never_public(org):
    assert "cost_estimate" not in public_spec(parse_workbook())
    assert "cost_estimate" not in org.get("/spec").json()
    assert org.get("/spec/internal/cost-estimate").status_code == 403


def test_every_check_has_a_detector_and_every_feature_is_mapped():
    s = parse_workbook()
    assert {c["code"] for c in s["audit_checklist"]["checks"]} <= set(REG)
    cov = traceability.coverage(s)
    assert cov["features_unmapped"] == []
    assert cov["integrations_unmapped"] == []
    assert cov["modules"] == 76


def test_diff_detects_changes_and_missing_sheet_rejected():
    raw = DEFAULT_WORKBOOK.read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb["Module Catalog"]
    ws.append([77, "SEO", "Log File Analyser", "Bot crawl analysis", "P2", "M", 2])
    ws2 = wb["Pricing Matrix"]
    for row in ws2.iter_rows():
        if row[0].value == "Rank Tracker keywords":
            row[1].value = "300 / mo"
    buf = io.BytesIO()
    wb.save(buf)
    d = diff_specs(parse_workbook(raw), parse_workbook(buf.getvalue()))
    sec = {x["section"]: x for x in d["sections"]}
    assert [m["name"] for m in sec["Modules"]["added"]] == ["Log File Analyser"]
    assert sec["Pricing"]["changed"][0]["key"] == "Rank Tracker keywords"
    del wb["Roadmap"]
    buf = io.BytesIO()
    wb.save(buf)
    try:
        parse_workbook(buf.getvalue())
        raise AssertionError("expected SpecError")
    except ValueError as e:
        assert "Roadmap" in str(e)


def test_refresh_requires_review_by_platform_admin(client, org):
    from tests.conftest import Org
    assert org.post("/spec/refresh/upload", files={"file": ("p.xlsx", b"x")}).status_code == 403
    admin = Org(client, email="admin@genranq-platform.io")
    from app.admin import set_pa
    from app.db import system_session
    from app import models as m
    with system_session() as db:
        slug = db.get(m.Organisation, admin.org_id).slug
    set_pa(admin.email, slug, True)
    admin.tokens = client.post("/api/auth/login", json={"email": admin.email, "password": admin.password, "organisation": slug}).json()
    raw = DEFAULT_WORKBOOK.read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    wb["Audit Checklist"].insert_rows(5)
    wb["Audit Checklist"]["A5"], wb["Audit Checklist"]["B5"], wb["Audit Checklist"]["C5"] = "CHK-99", "Setup", "Favicon present"
    wb["Audit Checklist"]["D5"], wb["Audit Checklist"]["E5"], wb["Audit Checklist"]["F5"], wb["Audit Checklist"]["G5"] = "Crawler", "Y", "Low", "Web Dev"
    buf = io.BytesIO()
    wb.save(buf)
    r = admin.post("/spec/refresh/upload", files={"file": ("plan.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "pending_review"
    added = {x["section"]: x for x in body["diff"]["sections"]}["Audit checks"]["added"]
    assert added[0]["code"] == "CHK-99"
    # nothing applied before approval
    assert len(admin.get("/checklist").json()) == 62
    assert admin.post(f"/spec/snapshots/{body['snapshot_id']}/reject").json()["status"] == "rejected"
    assert len(admin.get("/checklist").json()) == 62
