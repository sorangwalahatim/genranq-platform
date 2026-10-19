"""Parse the Genranq Build Plan workbook (the product source of truth) into a
normalised dict. Used to seed the platform, build navigation/traceability and
to diff a refreshed workbook against the currently-loaded specification.

The workbook is opened read-only and never written back.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

import openpyxl

REQUIRED_SHEETS = [
    "README", "Module Catalog", "Feature Backlog", "Database Schema", "API Integrations",
    "Roadmap", "Audit Checklist", "Pricing Matrix", "Cost Estimate",
]

DEFAULT_WORKBOOK = Path(__file__).with_name("Genranq_Build_Plan.xlsx")


class SpecError(ValueError):
    pass


def _rows(ws) -> list[list[Any]]:
    out = []
    for r in ws.iter_rows(values_only=True):
        vals = [c.strip() if isinstance(c, str) else c for c in r]
        if any(v not in (None, "") for v in vals):
            out.append(vals)
    return out


def _after_header(rows, first_header: str) -> list[list[Any]]:
    for i, r in enumerate(rows):
        if r and str(r[0] or "").strip() == first_header:
            return rows[i + 1:]
    raise SpecError(f"Header row starting with '{first_header}' not found")


def _num(v):
    if v is None or v == "":
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def _split_cols(col: str) -> list[str]:
    """'name/email/phone' -> ['name','email','phone']; 'country/lang' -> ['country','lang']."""
    return [c.strip() for c in col.split("/") if c.strip()]


def parse_readme(rows):
    items = []
    for r in rows:
        if len(r) >= 3 and r[1] and r[2] and r[1] != "Sheet":
            items.append({"key": r[1], "value": r[2]})
    title = next((r[1] for r in rows if r[1] and "GENRANQ" in str(r[1]).upper()), "GENRANQ")
    return {"title": title, "items": items}


def parse_modules(rows):
    out = []
    for r in _after_header(rows, "#"):
        if r[0] is None or r[2] is None:
            continue
        out.append({
            "num": int(float(r[0])), "area": r[1], "name": r[2], "description": r[3],
            "priority": r[4], "complexity": r[5], "phase": int(float(r[6])),
        })
    return out


def parse_features(rows):
    out = []
    for r in _after_header(rows, "ID"):
        if not r[0] or not str(r[0]).startswith("F-"):
            continue
        out.append({"id": r[0], "epic": r[1], "name": r[2], "moscow": r[3],
                    "points": _num(r[4]), "acceptance": r[5]})
    return out


def parse_schema(rows):
    tables: dict[str, dict] = {}
    for r in _after_header(rows, "Table"):
        table, col, typ, key, notes = (list(r) + [None] * 5)[:5]
        if not table or not col:
            continue
        t = tables.setdefault(table, {"name": table, "columns": [], "time_series": False})
        if notes and "ClickHouse" in str(notes):
            t["time_series"] = True
        t["columns"].append({"spec_column": col, "columns": _split_cols(col), "type": typ,
                             "key": key or "", "notes": notes or ""})
    return list(tables.values())


def parse_integrations(rows):
    out = []
    for r in _after_header(rows, "Provider"):
        if not r[0]:
            continue
        out.append({"provider": r[0], "category": r[1], "purpose": r[2], "auth": r[3],
                    "quota": r[4], "phase": int(float(r[5])) if r[5] is not None else None})
    return out


def parse_roadmap(rows):
    phases, totals = [], {}
    for r in _after_header(rows, "Phase"):
        if r[0] and re.match(r"^\d+\.", str(r[0])):
            phases.append({"phase": int(str(r[0]).split(".")[0]), "name": r[0], "focus": r[1],
                           "deliverables": r[2], "squad": r[3], "depends_on": r[4],
                           "weeks": _num(r[5])})
        elif r[4] and r[5] is not None:
            totals[str(r[4])] = _num(r[5])
    return {"phases": phases, "totals": totals}


def parse_checks(rows):
    out, note = [], None
    for r in _after_header(rows, "Code"):
        if r[0] and str(r[0]).startswith("CHK-"):
            out.append({"code": r[0], "category": r[1], "check": r[2], "data_source": r[3],
                        "auto": r[4], "severity": r[5], "owner": r[6]})
        elif r[0]:
            note = r[0]
    return {"checks": out, "note": note}


def parse_pricing(rows):
    body = _after_header(rows, "Feature")
    header = next(r for r in rows if r[0] == "Feature")
    tiers = [h for h in header[1:] if h]
    feats = []
    for r in body:
        if not r[0]:
            continue
        feats.append({"feature": r[0], "values": {t: (str(r[i + 1]) if r[i + 1] is not None else "—")
                                                  for i, t in enumerate(tiers)}})
    return {"tiers": tiers, "features": feats}


def parse_cost(rows):
    """Internal only — never exposed to customer-facing endpoints."""
    assumptions, phases, summary = {}, [], {}
    for r in rows:
        k = str(r[0] or "")
        if k.startswith("Hours / dev") or k.startswith("Blended rate") or k.startswith("Contingency"):
            assumptions[k] = _num(r[1])
        elif re.match(r"^\d+\.", k) or k.startswith("QA /"):
            phases.append({"phase": k, "weeks": _num(r[1]), "devs": _num(r[2]),
                           "dev_weeks": _num(r[3]), "hours": _num(r[4]), "cost": _num(r[5])})
        elif k.startswith("TOTAL"):
            summary["total"] = {"dev_weeks": _num(r[3]), "hours": _num(r[4]), "cost": _num(r[5])}
        elif k.startswith("GRAND TOTAL"):
            summary["grand_total"] = _num(r[5])
    return {"assumptions": assumptions, "phases": phases, "summary": summary}


def parse_workbook(source: str | Path | bytes | None = None) -> dict:
    if source is None:
        source = DEFAULT_WORKBOOK
    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True, read_only=True)
        raw = bytes(source)
    else:
        raw = Path(source).read_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise SpecError(f"Workbook is missing required sheets: {', '.join(missing)}")
    S = {name: _rows(wb[name]) for name in REQUIRED_SHEETS}
    spec = {
        "readme": parse_readme(S["README"]),
        "modules": parse_modules(S["Module Catalog"]),
        "features": parse_features(S["Feature Backlog"]),
        "schema": parse_schema(S["Database Schema"]),
        "integrations": parse_integrations(S["API Integrations"]),
        "roadmap": parse_roadmap(S["Roadmap"]),
        "audit_checklist": parse_checks(S["Audit Checklist"]),
        "pricing": parse_pricing(S["Pricing Matrix"]),
        "cost_estimate": parse_cost(S["Cost Estimate"]),
    }
    spec["meta"] = {
        "sheets_read": REQUIRED_SHEETS,
        "sha256": hashlib.sha256(raw).hexdigest(),
        "counts": {
            "modules": len(spec["modules"]), "features": len(spec["features"]),
            "tables": len(spec["schema"]), "integrations": len(spec["integrations"]),
            "phases": len(spec["roadmap"]["phases"]),
            "checks": len(spec["audit_checklist"]["checks"]),
            "pricing_rows": len(spec["pricing"]["features"]),
        },
    }
    return spec


def public_spec(spec: dict) -> dict:
    """Spec with internal cost data stripped (safe for any authenticated user)."""
    return {k: v for k, v in spec.items() if k != "cost_estimate"}


# ---------------------------------------------------------------- diffing
def _index(items, key):
    return {str(i[key]): i for i in items}


def _diff_list(old, new, key, label):
    o, n = _index(old, key), _index(new, key)
    added = [n[k] for k in n if k not in o]
    removed = [o[k] for k in o if k not in n]
    changed = []
    for k in n:
        if k in o and json.dumps(o[k], sort_keys=True, default=str) != json.dumps(n[k], sort_keys=True, default=str):
            fields = sorted(f for f in set(o[k]) | set(n[k]) if o[k].get(f) != n[k].get(f))
            changed.append({"key": k, "fields": fields,
                            "before": {f: o[k].get(f) for f in fields},
                            "after": {f: n[k].get(f) for f in fields}})
    return {"section": label, "added": added, "removed": removed, "changed": changed}


def diff_specs(old: dict, new: dict) -> dict:
    sections = [
        _diff_list(old["modules"], new["modules"], "name", "Modules"),
        _diff_list(old["features"], new["features"], "id", "Features"),
        _diff_list(old["schema"], new["schema"], "name", "Database tables"),
        _diff_list(old["integrations"], new["integrations"], "provider", "Integrations"),
        _diff_list(old["roadmap"]["phases"], new["roadmap"]["phases"], "phase", "Roadmap phases"),
        _diff_list(old["audit_checklist"]["checks"], new["audit_checklist"]["checks"], "code", "Audit checks"),
        _diff_list(old["pricing"]["features"], new["pricing"]["features"], "feature", "Pricing"),
    ]
    total = sum(len(s["added"]) + len(s["removed"]) + len(s["changed"]) for s in sections)
    return {"total_changes": total, "sections": sections,
            "old_sha256": old.get("meta", {}).get("sha256"), "new_sha256": new.get("meta", {}).get("sha256")}


if __name__ == "__main__":
    s = parse_workbook()
    print(json.dumps(s["meta"], indent=2))
