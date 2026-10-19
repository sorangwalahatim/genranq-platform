"""Seed data for the checklists module, parsed from the user's Excel files.

  * Pre___Post_Launch_Audit.xlsx -> "Pre Launch" and "Post Launch" sheets
    (columns: Task, Details, Owner, Deliverable?, Acceptance Criteria, Notes, Status, Priority)
  * SEO_Tech_Audit_Checklist.xlsx -> "Audit Summary" sheet (24 items) and
    "SEO-Tech-Audit-Checklist" sheet (graded checks: Check, Grade, Observations/Notes,
    Responsibility, Status)

Returns plain dicts; the router turns them into rows (idempotent per project).
"""
from __future__ import annotations

PRE_POST_XLSX = "/home/hatch/workspace/user/files/Pre___Post_Launch_Audit.xlsx"
TECH_AUDIT_XLSX = "/home/hatch/workspace/user/files/SEO_Tech_Audit_Checklist.xlsx"

KINDS = ("pre_launch", "post_launch", "tech_audit")


def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise RuntimeError("openpyxl is not installed — checklist seeding needs it (pip install openpyxl)")
    return openpyxl


def _norm_status(raw: str | None) -> str:
    s = (raw or "").strip().lower()
    if s in ("pass", "passed", "done", "complete", "completed", "yes", "ok"):
        return "pass"
    if s in ("fail", "failed", "no"):
        return "fail"
    if s in ("partial", "partly"):
        return "partial"
    return "pending"


def _norm_owner(raw: str | None) -> str | None:
    s = (raw or "").strip()
    if not s:
        return None
    low = s.lower()
    if "developer" in low or "dev " in low or low == "dev":
        return "Developer"
    if "content" in low or "writer" in low:
        return "Content Writer"
    if "seo" in low or "analyst" in low or "analytics" in low:
        return "SEO Lead"
    if "design" in low:
        return "Designer"
    return s[:60]


def _cell(v) -> str:
    return str(v).strip() if v is not None else ""


def _parse_pre_post() -> dict[str, list[dict]]:
    """Section headers (Task set but Owner/Details/Acceptance all empty) are skipped."""
    openpyxl = _need_openpyxl()
    wb = openpyxl.load_workbook(PRE_POST_XLSX, data_only=True, read_only=True)
    out: dict[str, list[dict]] = {"pre_launch": [], "post_launch": []}
    for sheet_name, kind in (("Pre Launch", "pre_launch"), ("Post Launch", "post_launch")):
        ws = wb[sheet_name]
        order = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            cells = list(row) + [None] * 8  # read_only rows can be short
            task, details, owner = _cell(cells[0]), _cell(cells[1]), _cell(cells[2])
            acceptance, notes = _cell(cells[4]), _cell(cells[5])
            priority = _cell(cells[7])
            if not task:
                continue
            if not (owner or details or acceptance):
                continue  # section header like "Tracking" / "Technical"
            detail_parts = [p for p in (details, f"Acceptance: {acceptance}" if acceptance else "", notes) if p]
            order += 1
            out[kind].append({
                "title": task[:300],
                "detail": "\n".join(detail_parts) or None,
                "owner_role": _norm_owner(owner),
                "priority": priority.lower() if priority.lower() in ("low", "medium", "high", "critical") else None,
                "status": "pending",
                "sort_order": order,
            })
    wb.close()
    return out


def _parse_tech_audit() -> list[dict]:
    openpyxl = _need_openpyxl()
    wb = openpyxl.load_workbook(TECH_AUDIT_XLSX, data_only=True, read_only=True)
    items: list[dict] = []
    order = 0
    # 1) Audit Summary sheet: rows 3.. : No | Check List | Remarks/Update
    ws = wb["Audit Summary"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        cells = list(row) + [None] * 3
        title, remarks = _cell(cells[1]), _cell(cells[2])
        if not title:
            continue
        order += 1
        items.append({"title": title[:300], "detail": remarks or None, "owner_role": None,
                      "priority": None, "status": "pending", "sort_order": order})
    # 2) Detailed graded checklist: row2 header Check | Grade | Observations/Notes | Responsibility | Status
    ws2 = wb["SEO-Tech-Audit-Checklist"]
    for row in ws2.iter_rows(min_row=3, values_only=True):
        cells = list(row) + [None] * 5
        check, grade, notes, resp, status = _cell(cells[0]), _cell(cells[1]), _cell(cells[2]), _cell(cells[3]), _cell(cells[4])
        if not check:
            continue
        order += 1
        detail_parts = [p for p in ((f"Grade in sheet: {grade}" if grade else ""), notes) if p]
        items.append({"title": check[:300], "detail": "\n".join(detail_parts) or None,
                      "owner_role": _norm_owner(resp), "priority": None,
                      "status": _norm_status(grade or status), "sort_order": order})
    wb.close()
    return items


def load_seed_data() -> dict[str, list[dict]]:
    """{kind: [ {title, detail, owner_role, priority, status, sort_order}, ... ]}."""
    data = _parse_pre_post()
    data["tech_audit"] = _parse_tech_audit()
    return data
