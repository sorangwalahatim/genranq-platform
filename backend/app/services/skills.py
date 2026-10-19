"""Prompt skills: the user's analysis prompts as project-wise, repeatable skills.

Each skill = a prompt template + the evidence it needs + a deterministic analysis.

  1. inputs   — filled from the Project profile (industry, product, audience, USP, tone,
                keywords, competitors, location…) and overridable per run
  2. evidence — collected for real: client and competitor pages are fetched and measured
                (title/meta lengths, H1–H6, schema, CTAs, FAQs, trust signals, word count,
                images, tables, videos, forms, tone metrics), plus Search Console queries
  3. analysis — every skill produces a data analysis from that evidence without AI; when an AI
                provider is connected the full prompt runs on the evidence and its answer
                (narrative + tables) is stored alongside
  4. output   — stored per project (SkillRun) with the exact prompt used, exportable to Excel,
                schedulable as an activity ("skill:<key>") for a person, a team or an AI agent

Nothing is invented: rows produced by templates rather than data are labelled as such.
"""
from __future__ import annotations

import io
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from sqlalchemy import func, select

from .. import models as m
from .. import models_monitor as mm
from . import keywords as kw

YEAR = date.today().year

PROFILE_FIELDS = [
    ("website", "Website URL"), ("brand", "Brand / business name"), ("industry", "Industry / niche"), ("product", "Product / service"),
    ("stage", "Business stage"), ("revenue_model", "Revenue model"), ("audience", "Target audience"), ("persona", "Primary persona"),
    ("location", "Location / service area"), ("usp", "Unique selling proposition"), ("brand_promise", "Brand goal / promise"),
    ("tone", "Desired tone"), ("primary_keywords", "Primary keywords"), ("secondary_keywords", "Secondary keywords"),
    ("long_tail_keywords", "Long-tail keywords"), ("competitors", "Competitor URLs"), ("pain_points", "Customer pain points"),
    ("proof_points", "Proof points / results (metrics)"), ("budget_note", "Typical customer budget"),
]

CLICHES = ["leverage", "synergy", "synergistic", "solutions", "cutting-edge", "world-class", "best-in-class", "seamless", "robust",
           "innovative", "state-of-the-art", "game-changer", "next-generation", "one-stop", "holistic", "empower", "unlock", "revolutionize",
           "utilize", "paradigm", "turnkey", "scalable", "bespoke", "end-to-end", "in today's fast-paced"]
CTA_WORDS = r"\b(get|start|book|buy|try|request|contact|call|schedule|download|sign ?up|subscribe|join|order|claim|talk|see|learn|explore|shop|quote|demo|free)\b"
TRUST_RX = {
    "testimonials": r"testimonial|what (our )?(clients|customers) say|review(s|ed)? by",
    "ratings": r"\b\d(\.\d)?\s*(/\s*5|out of 5|stars?)\b|aggregaterating",
    "certifications": r"certified|accredited|licen[cs]ed|iso \d|registered|member of",
    "awards": r"award|winner|recogni[sz]ed",
    "guarantee": r"guarantee|money[- ]back|warranty|risk[- ]free",
    "experience": r"\b\d+\+?\s*(years|yrs)\b|since (19|20)\d\d",
    "social_proof": r"\b\d[\d,]*\+?\s*(clients|customers|users|businesses|brands|projects|reviews)\b|trusted by",
    "case_studies": r"case stud(y|ies)|success stor(y|ies)",
    "privacy_security": r"privacy policy|secure|ssl|gdpr|encrypted",
}
YMYL_RX = r"\b(health|medical|doctor|clinic|dental|surgery|medicine|drug|symptom|finance|loan|mortgage|invest|insurance|tax|legal|lawyer|attorney|visa|immigration|pension|credit|bank|safety|pharma|therapy|diet)\b"
STOP = kw.STOP | set("""also more most other some such only own same so too very just will one all any each few both but if then there their they them
we us our its it's his her he she not no nor been being have has had having would could may might must shall get got make made use used using
via per new like well see here now get page site click read home menu contact privacy policy cookie cookies rights reserved copyright""".split())


def utcnow():
    return datetime.now(timezone.utc)


# ------------------------------------------------------------------ evidence: page facts
def page_facts(url: str) -> dict:
    """Fetch one page (SSRF-guarded crawler) and measure everything the prompts ask about."""
    from .crawler import Crawler, CrawlBlocked
    if not url:
        return {"url": url, "error": "no URL"}
    if "://" not in url:
        url = "https://" + url.strip("/") + "/"
    c = Crawler(url, max_pages=1, render_sample=0, mobile_sample=0, external_check_limit=0, delay=0)
    try:
        r = c.fetch(url)
    except CrawlBlocked as e:
        return {"url": url, "error": str(e)}
    except Exception as e:  # network errors are evidence too
        return {"url": url, "error": f"{type(e).__name__}: {e}"[:200]}
    p = c.parse(url, r)
    out = {"url": url, "status": p.get("status"), "error": p.get("error")}
    body = r.get("body") or ""
    if p.get("status") != 200 or not body:
        return out
    soup = BeautifulSoup(body, "lxml")
    for t in soup(["script", "style", "noscript", "svg"]):
        t.decompose()
    title = (soup.title.get_text(" ", strip=True) if soup.title else "") or ""
    md = soup.find("meta", attrs={"name": re.compile("^description$", re.I)})
    meta = (md.get("content") or "").strip() if md else ""
    heads = {f"h{i}": [h.get_text(" ", strip=True)[:200] for h in soup.find_all(f"h{i}")] for i in range(1, 7)}
    order = [int(h.name[1]) for h in soup.find_all(re.compile("^h[1-6]$"))]
    skips = sum(1 for a, b in zip(order, order[1:]) if b - a > 1)
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = re.sub(r"\s+", " ", main.get_text(" ", strip=True))
    full_text = re.sub(r"\s+", " ", (soup.body or soup).get_text(" ", strip=True))
    words = re.findall(r"[A-Za-z][A-Za-z'\-]+", text)
    paras = [re.sub(r"\s+", " ", x.get_text(" ", strip=True)) for x in main.find_all("p")]
    paras = [x for x in paras if len(x.split()) >= 8]
    ctas = []
    for el in soup.find_all(["a", "button", "input"]):
        label = (el.get("value") if el.name == "input" else el.get_text(" ", strip=True)) or ""
        cls = " ".join(el.get("class") or []) if isinstance(el.get("class"), list) else (el.get("class") or "")
        if el.name == "input" and (el.get("type") or "").lower() not in ("submit", "button"):
            continue
        if label and len(label) <= 40 and (el.name != "a" or re.search(r"btn|button|cta", cls, re.I) or re.search(CTA_WORDS, label, re.I)):
            if re.search(CTA_WORDS, label, re.I) or el.name != "a":
                ctas.append(label)
    forms = soup.find_all("form")
    form_fields = [len(f.find_all(["input", "select", "textarea"])) for f in forms]
    faq_q = [h for i in range(2, 5) for h in heads[f"h{i}"] if h.endswith("?")]
    faq_q += [s.get_text(" ", strip=True)[:200] for s in soup.find_all("summary")]
    videos = len(soup.find_all("video")) + sum(1 for f in soup.find_all("iframe") if re.search(r"youtube|vimeo|wistia|loom", f.get("src") or ""))
    low = full_text.lower() + " " + body.lower()[:200000]
    trust = {k: bool(re.search(rx, low)) for k, rx in TRUST_RX.items()}
    trust["author"] = bool(p.get("author"))
    trust["phone"] = bool(p.get("phones"))
    trust["address"] = bool(re.search(r"postaladdress|streetaddress|\b\d{1,5} [a-z]+ (street|st|road|rd|avenue|ave|lane|ln)\b", low))
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", " ".join(paras)) if len(s.split()) >= 3]
    sl = [len(s.split()) for s in sentences]
    passive = sum(1 for s in sentences if re.search(r"\b(is|are|was|were|be|been|being)\s+\w+ed\b", s.lower()))
    syll = sum(_syllables(w) for w in words) or 1
    flesch = round(206.835 - 1.015 * (len(words) / max(1, len(sentences))) - 84.6 * (syll / max(1, len(words))), 1) if sentences else None
    lw = [w.lower() for w in words]
    hero = ""
    h1 = soup.find("h1")
    if h1:
        nxt = h1.find_next(["p", "h2"])
        hero = (h1.get_text(" ", strip=True) + " — " + (nxt.get_text(" ", strip=True) if nxt else ""))[:300]
    usp = [s for s in sentences if re.search(r"\bonly\b|#1|number one|leading|trusted by|\bguarantee|\bfaster|\bcheaper|\bunlike\b|\d+%|award", s.lower())][:8]
    out.update({
        "title": title, "title_length": len(title), "meta_description": meta, "meta_length": len(meta),
        "canonical": p.get("canonical"), "noindex": bool(p.get("noindex")), "lang": p.get("lang"), "viewport": p.get("viewport"),
        "headings": heads, "heading_counts": {k: len(v) for k, v in heads.items()}, "heading_level_skips": skips,
        "schema_types": p.get("schema_types") or [], "schema_errors": p.get("schema_errors") or [],
        "ctas": list(dict.fromkeys(ctas))[:20], "cta_count": len(ctas), "forms": len(forms), "form_fields": form_fields,
        "faq_questions": list(dict.fromkeys(faq_q))[:30], "faq_schema": "FAQPage" in (p.get("schema_types") or []),
        "trust_signals": trust, "trust_score": sum(1 for v in trust.values() if v),
        "word_count": len(words), "images": p.get("images", 0), "images_missing_alt": len(p.get("images_missing_alt") or []),
        "tables": p.get("tables", 0), "lists": p.get("lists", 0), "videos": videos,
        "internal_links": len(p.get("internal_links") or []) if isinstance(p.get("internal_links"), list) else p.get("internal_links"),
        "external_links": len(p.get("external_links") or []) if isinstance(p.get("external_links"), list) else p.get("external_links"),
        "date_modified": p.get("date_modified"), "hero": hero, "usp_statements": usp,
        "tone": {"avg_sentence_words": round(sum(sl) / len(sl), 1) if sl else None,
                 "long_sentences_pct": round(100 * sum(1 for x in sl if x > 25) / len(sl)) if sl else None,
                 "passive_pct": round(100 * passive / len(sentences)) if sentences else None, "flesch_reading_ease": flesch,
                 "you_per_100w": round(100 * sum(lw.count(x) for x in ("you", "your", "you're")) / max(1, len(lw)), 2),
                 "we_per_100w": round(100 * sum(lw.count(x) for x in ("we", "our", "us", "we're")) / max(1, len(lw)), 2),
                 "exclamations": text.count("!"), "questions": text.count("?"),
                 "cliches": sorted({c for c in CLICHES if c in text.lower()})},
        "sample_paragraphs": paras[:6], "long_sentences": [s for s in sentences if len(s.split()) > 30][:3],
        "text": text[:12000], "ymyl": bool(re.search(YMYL_RX, (title + " " + text[:5000]).lower())),
        "definitions": len(re.findall(r"\b(is a|is an|refers to|means)\b", text.lower())),
    })
    return out


def _syllables(w: str) -> int:
    w = w.lower()
    n = len(re.findall(r"[aeiouy]+", w))
    if w.endswith("e") and n > 1:
        n -= 1
    return max(1, n)


# ------------------------------------------------------------------ profile + inputs
def get_profile(db, project) -> dict:
    p = db.get(mm.ProjectProfile, project.id)
    data = dict(p.data) if p else {}
    if not data.get("website"):
        site = db.scalar(select(m.Site).where(m.Site.project_id == project.id).limit(1))
        if site:
            data["website"] = site.domain if "://" in site.domain else "https://" + site.domain.strip("/") + "/"
    if not data.get("competitors"):
        data["competitors"] = ["https://" + c.domain.split("://")[-1].strip("/") + "/" for c in
                               db.scalars(select(m.Competitor).where(m.Competitor.project_id == project.id).limit(10))]
    if not data.get("brand"):
        data["brand"] = project.client_name or project.name
    return data


def _list(v) -> list:
    if not v:
        return []
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    return [x.strip() for x in re.split(r"[\n,;]+", str(v)) if x.strip()]


def merged_inputs(profile: dict, inputs: dict) -> dict:
    v = {**profile, **{k: x for k, x in (inputs or {}).items() if x not in (None, "", [])}}
    for k in ("primary_keywords", "secondary_keywords", "long_tail_keywords", "competitors", "competitor_urls", "pain_points"):
        if k in v:
            v[k] = _list(v[k])
    v.setdefault("url", v.get("website"))
    v["competitor_urls"] = _list(v.get("competitor_urls")) or _list(v.get("competitors"))
    pk = _list(v.get("primary_keywords"))
    v.setdefault("primary_query", (v.get("query") or (pk[0] if pk else None)))
    v.setdefault("year", YEAR)
    return v


# ------------------------------------------------------------------ evidence collectors
def collect(db, project, needs: list[str], v: dict) -> dict:
    ev: dict = {}
    if "client_page" in needs:
        ev["client_page"] = page_facts(v.get("url"))
    if "competitor_pages" in needs:
        lim = int(v.get("max_competitors") or 10)
        ev["competitor_pages"] = [page_facts(u) for u in v["competitor_urls"][:lim]]
    if "gsc_queries" in needs:
        ev["gsc_queries"] = gsc_queries(db, project)
    if "keywords" in needs:
        ev["keywords"] = [{"keyword": k.phrase, "volume": k.search_volume, "difficulty": k.difficulty, "target_url": k.target_url}
                          for k in db.scalars(select(m.Keyword).where(m.Keyword.project_id == project.id).limit(500))]
        ev["competitor_keywords"] = [{"domain": c.domain, "keyword": c.keyword, "position": c.position, "volume": c.volume, "own": c.is_own}
                                     for c in db.scalars(select(mm.CompetitorKeyword).where(mm.CompetitorKeyword.project_id == project.id).limit(2000))]
    return ev


def gsc_queries(db, project, days=90, limit=500) -> list[dict]:
    anchor = db.scalar(select(func.max(mm.SearchMetric.date)).where(mm.SearchMetric.project_id == project.id))
    if not anchor:
        return []
    rows = db.execute(select(mm.SearchMetric.query, func.sum(mm.SearchMetric.clicks), func.sum(mm.SearchMetric.impressions),
                             func.sum(mm.SearchMetric.position * mm.SearchMetric.impressions))
                      .where(mm.SearchMetric.project_id == project.id, mm.SearchMetric.query.is_not(None), mm.SearchMetric.grain == "detail",
                             mm.SearchMetric.date > anchor - timedelta(days=days))
                      .group_by(mm.SearchMetric.query).order_by(func.sum(mm.SearchMetric.impressions).desc()).limit(limit)).all()
    return [{"query": q, "clicks": int(c or 0), "impressions": int(i or 0), "position": round(float(wp) / i, 1) if i and wp is not None else None}
            for q, c, i, wp in rows]


def _ok_pages(ev, key="competitor_pages"):
    return [p for p in ev.get(key) or [] if p.get("status") == 200]


def _host(u):
    return urlparse(u if "://" in (u or "") else "https://" + (u or "")).netloc.replace("www.", "")


def _need_pages(ev, competitors=True):
    cp = ev.get("client_page") or {}
    msg = []
    if cp and cp.get("status") != 200:
        msg.append(f"Client page could not be fetched ({cp.get('url')}: status {cp.get('status')} {cp.get('error') or ''}).")
    if competitors and not _ok_pages(ev):
        msg.append("No competitor page could be fetched — add competitor URLs in the project profile or in the run inputs.")
    return msg


def _kw_in(text, kws):
    t = (text or "").lower()
    return any(k.lower() in t for k in kws) if kws else None


def _median(xs):
    xs = sorted(x for x in xs if x is not None)
    return xs[len(xs) // 2] if xs else None


# ------------------------------------------------------------------ deterministic analyses (one per skill)
def a_meta_gap(v, ev) -> dict:
    cp = ev.get("client_page") or {}
    comps = _ok_pages(ev)
    kws = v.get("primary_keywords") or ([v["primary_query"]] if v.get("primary_query") else [])
    sites = ([cp] if cp.get("status") == 200 else []) + comps
    rows = []
    for i, p in enumerate(sites):
        rows.append({"Site": "Our site" if (i == 0 and cp.get("status") == 200) else _host(p["url"]), "URL": p["url"],
                     "Title": p["title"], "Title length": p["title_length"],
                     "Title length verdict": "ok" if 30 <= p["title_length"] <= 60 else ("too long" if p["title_length"] > 60 else "too short"),
                     "Keyword in title": _yn(_kw_in(p["title"], kws)), "Meta description": p["meta_description"], "Meta length": p["meta_length"],
                     "Meta length verdict": "ok" if 70 <= p["meta_length"] <= 160 else ("missing" if not p["meta_length"] else ("too long" if p["meta_length"] > 160 else "too short")),
                     "Keyword in meta": _yn(_kw_in(p["meta_description"], kws)), "CTA in meta": _yn(bool(re.search(CTA_WORDS, p["meta_description"].lower()))),
                     "Number/proof in title or meta": _yn(bool(re.search(r"\d", p["title"] + p["meta_description"]))),
                     "USP statements": " | ".join(p["usp_statements"][:2]), "H1": " | ".join(p["headings"]["h1"][:2]),
                     "H2 count": p["heading_counts"]["h2"], "Heading level skips": p["heading_level_skips"],
                     "Schema": ", ".join(p["schema_types"]) or "none", "Word count": p["word_count"], "Trust signals": p["trust_score"],
                     "CTAs": ", ".join(p["ctas"][:5]), "Forms": p["forms"], "FAQ questions": len(p["faq_questions"]),
                     "Images": p["images"], "Videos": p["videos"], "Indexable": _yn(not p["noindex"]),
                     "Readability (Flesch)": p["tone"]["flesch_reading_ease"]})
    gaps, actions = [], []
    if cp.get("status") == 200 and comps:
        def cmp(area, ours, theirs, better, action, fmt=lambda x: x):
            med = _median(theirs)
            if med is None or ours is None:
                return
            behind = (ours < med) if better == "higher" else (ours > med)
            gaps.append({"Area": area, "Our site": fmt(ours), "Competitor median": fmt(med), "Gap": "behind" if behind else "on par / ahead",
                         "Action": action if behind else "Keep"})
            if behind:
                actions.append(action)
        t = cp["title_length"]
        if not (30 <= t <= 60):
            actions.insert(0, f"Rewrite the title to 50–60 characters (now {t}) with the primary keyword near the start and a differentiator.")
        if not cp["meta_length"] or cp["meta_length"] > 160 or cp["meta_length"] < 70:
            actions.insert(0, f"Rewrite the meta description to 140–155 characters (now {cp['meta_length']}) with benefit + proof + CTA.")
        if kws and not _kw_in(cp["title"], kws):
            actions.insert(0, f"Add the primary keyword “{kws[0]}” to the title.")
        cmp("Content depth (words)", cp["word_count"], [c["word_count"] for c in comps], "higher", "Expand the page to cover the subtopics competitors cover (see heading gaps).")
        cmp("H2 sections", cp["heading_counts"]["h2"], [c["heading_counts"]["h2"] for c in comps], "higher", "Add H2 sections for the topics competitors cover.")
        cmp("Trust signals (of 12)", cp["trust_score"], [c["trust_score"] for c in comps], "higher", "Add proof: testimonials, ratings, certifications, years in business, case studies.")
        cmp("FAQ questions", len(cp["faq_questions"]), [len(c["faq_questions"]) for c in comps], "higher", "Add an FAQ block with FAQPage schema from real search questions.")
        cmp("Schema types", len(cp["schema_types"]), [len(c["schema_types"]) for c in comps], "higher", "Add structured data (Organization, Service/Product, FAQPage, Review where valid).")
        cmp("CTAs", cp["cta_count"], [c["cta_count"] for c in comps], "higher", "Add clear primary and secondary CTAs above the fold and after key sections.")
        cmp("Images", cp["images"], [c["images"] for c in comps], "higher", "Add original visuals (process, results, team) with descriptive alt text.")
        cmp("Readability (Flesch, higher = easier)", cp["tone"]["flesch_reading_ease"], [c["tone"]["flesch_reading_ease"] for c in comps], "higher",
            "Simplify sentences: shorter sentences, plain words, active voice.")
        if cp["noindex"]:
            actions.insert(0, "Remove the noindex directive — the page cannot rank.")
        if cp["heading_level_skips"]:
            actions.append(f"Fix {cp['heading_level_skips']} skipped heading level(s) so H2→H3→H4 nest correctly.")
    swot = _swot(cp, comps, gaps)
    return {"tables": {"Title & meta comparison": rows, "Gap analysis": gaps,
                       "Top 5 actions": [{"#": i + 1, "Action": a} for i, a in enumerate(list(dict.fromkeys(actions))[:5])],
                       "SWOT": swot}, "notes": _need_pages(ev)}


def _yn(b):
    return "—" if b is None else ("Yes" if b else "No")


def _swot(cp, comps, gaps):
    if not cp or cp.get("status") != 200:
        return []
    s = [{"Type": "Strength", "Point": f"{g['Area']}: {g['Our site']} vs median {g['Competitor median']}"} for g in gaps if g["Gap"] != "behind"]
    w = [{"Type": "Weakness", "Point": f"{g['Area']}: {g['Our site']} vs median {g['Competitor median']}"} for g in gaps if g["Gap"] == "behind"]
    o = []
    if comps and sum(1 for c in comps if c["faq_schema"]) <= len(comps) / 2:
        o.append({"Type": "Opportunity", "Point": "Most competitors lack FAQ schema — an answer-first FAQ can win rich results and AI citations."})
    if comps and _median([c["videos"] for c in comps]) == 0:
        o.append({"Type": "Opportunity", "Point": "Competitors don't use video — a short explainer video would differentiate."})
    if comps and _median([c["tone"]["flesch_reading_ease"] for c in comps]) and (_median([c["tone"]["flesch_reading_ease"] for c in comps]) or 0) < 50:
        o.append({"Type": "Opportunity", "Point": "Competitor copy is hard to read — clearer copy is an easy differentiator."})
    t = [{"Type": "Threat", "Point": f"{_host(c['url'])} shows {c['trust_score']} trust signals and {c['word_count']} words."}
         for c in comps if c["trust_score"] > cp["trust_score"] and c["word_count"] > cp["word_count"]][:3]
    return s + w + o + t


def a_competitor_pages(v, ev) -> dict:
    cp = ev.get("client_page") or {}
    comps = _ok_pages(ev)
    sites = ([cp] if cp.get("status") == 200 else []) + comps
    layout = []
    for p in sites:
        layout.append({"Site": _host(p["url"]), "URL": p["url"], "H1": " | ".join(p["headings"]["h1"][:1]),
                       "Sections (H2 flow)": " → ".join(p["headings"]["h2"][:12]), "H2": p["heading_counts"]["h2"], "H3": p["heading_counts"]["h3"],
                       "Word count": p["word_count"], "Images": p["images"], "Videos": p["videos"], "Tables": p["tables"], "Lists": p["lists"],
                       "FAQ": len(p["faq_questions"]), "Forms": p["forms"], "CTAs": ", ".join(p["ctas"][:4]),
                       "Trust signals": ", ".join(k for k, x in p["trust_signals"].items() if x), "Schema": ", ".join(p["schema_types"]),
                       "Updated date shown": _yn(p["date_modified"]), "Author shown": _yn(p["trust_signals"]["author"])})
    features = ["FAQ section", "FAQ schema", "Video", "Comparison table", "Testimonials", "Ratings", "Case studies", "Certifications", "Guarantee",
                "Pricing mentioned", "Author byline", "Last-updated date", "Lead form", "Phone number"]

    def has(p, f):
        ts = p["trust_signals"]
        return {"FAQ section": bool(p["faq_questions"]), "FAQ schema": p["faq_schema"], "Video": p["videos"] > 0, "Comparison table": p["tables"] > 0,
                "Testimonials": ts["testimonials"], "Ratings": ts["ratings"], "Case studies": ts["case_studies"], "Certifications": ts["certifications"],
                "Guarantee": ts["guarantee"], "Pricing mentioned": bool(re.search(r"\$|₹|€|£|pricing|price", p["text"].lower())),
                "Author byline": ts["author"], "Last-updated date": bool(p["date_modified"]), "Lead form": p["forms"] > 0, "Phone number": ts["phone"]}[f]
    matrix = [{"Feature": f, **{_host(p["url"]): ("✓" if has(p, f) else "—") for p in sites}} for f in features]
    gaps, recs = [], []
    if cp.get("status") == 200 and comps:
        for f in features:
            n = sum(1 for c in comps if has(c, f))
            if not has(cp, f) and n >= max(1, len(comps) / 2):
                gaps.append({"Gap": f, "Competitors with it": f"{n}/{len(comps)}", "Recommendation": f"Add {f.lower()} — most ranking competitors have it."})
        ours = {kw.stem(t) for h in cp["headings"]["h2"] for t in kw.tokens(h)}
        topic_count = Counter()
        examples = {}
        for c in comps:
            for h in c["headings"]["h2"]:
                for t in {kw.stem(t) for t in kw.tokens(h)}:
                    topic_count[t] += 1
                    examples.setdefault(t, h)
        for t, n in topic_count.most_common(40):
            if t not in ours and n >= max(2, len(comps) // 2):
                gaps.append({"Gap": f"Topic: {t}", "Competitors with it": f"{n}/{len(comps)}", "Recommendation": f"Cover it, e.g. a section like “{examples[t]}”."})
        med_w = _median([c["word_count"] for c in comps])
        if med_w and cp["word_count"] < 0.7 * med_w:
            recs.append(f"Content is {cp['word_count']} words vs a competitor median of {med_w}: expand depth on the missing topics.")
        recs += [g["Recommendation"] for g in gaps[:8]]
    return {"tables": {"Layout & quality signals": layout, "Comparison matrix": matrix, "Gaps": gaps,
                       "Recommendations": [{"#": i + 1, "Recommendation": r} for i, r in enumerate(dict.fromkeys(recs))]}, "notes": _need_pages(ev)}


def a_usp_heading(v, ev) -> dict:
    cp = ev.get("client_page") or {}
    comps = _ok_pages(ev)
    sites = ([cp] if cp.get("status") == 200 else []) + comps
    rows, heads = [], []
    for p in sites:
        h = p["headings"]
        rows.append({"Site": _host(p["url"]), "Hero (H1 + first line)": p["hero"], "USP statements found": " | ".join(p["usp_statements"][:3]) or "none found",
                     "USP in H1": _yn(bool(p["usp_statements"]) and any(any(w in h1.lower() for w in re.findall(r"[a-z]{5,}", p["usp_statements"][0].lower())[:4]) for h1 in h["h1"])),
                     "Tone (Flesch / you:we)": f"{p['tone']['flesch_reading_ease']} / {p['tone']['you_per_100w']}:{p['tone']['we_per_100w']}",
                     "Section flow (H2)": " → ".join(h["h2"][:15]), "CTAs": ", ".join(p["ctas"][:4])})
        heads.append({"Site": _host(p["url"]), "H1": " | ".join(h["h1"]) or "MISSING", "H1 count": len(h["h1"]), "H2 list": " | ".join(h["h2"][:20]),
                      "H3": len(h["h3"]), "H4": len(h["h4"]), "H5": len(h["h5"]), "H6": len(h["h6"]),
                      "Level skips": p["heading_level_skips"], "Question headings": len(p["faq_questions"]),
                      "Assessment": _heading_verdict(p, v)})
    gaps = []
    if cp.get("status") == 200 and comps:
        if not cp["usp_statements"]:
            gaps.append({"Aspect": "Primary USP clarity", "Gap & insight": "No differentiating claim (only/#1/guarantee/numbers) found on our page; competitors: "
                         + "; ".join(f"{_host(c['url'])}: {c['usp_statements'][0][:90]}" for c in comps if c["usp_statements"])[:400]})
        if len(cp["headings"]["h1"]) != 1:
            gaps.append({"Aspect": "H1", "Gap & insight": f"Our page has {len(cp['headings']['h1'])} H1 tags — use exactly one H1 carrying the main promise."})
        mh2 = _median([c["heading_counts"]["h2"] for c in comps])
        if mh2 and cp["heading_counts"]["h2"] < mh2:
            gaps.append({"Aspect": "H2 story", "Gap & insight": f"{cp['heading_counts']['h2']} H2s vs competitor median {mh2}: the page tells a shorter story."})
        if cp["heading_level_skips"]:
            gaps.append({"Aspect": "H3–H6 support", "Gap & insight": f"{cp['heading_level_skips']} skipped heading levels break the outline."})
    return {"tables": {"USP & messaging": rows, "Heading structure": heads, "Gaps": gaps}, "notes": _need_pages(ev)}


def _heading_verdict(p, v):
    issues = []
    h = p["headings"]
    if len(h["h1"]) != 1:
        issues.append(f"{len(h['h1'])} H1s")
    kws = v.get("primary_keywords") or []
    if kws and h["h1"] and not _kw_in(" ".join(h["h1"]), kws):
        issues.append("H1 lacks the primary keyword")
    if p["heading_level_skips"]:
        issues.append("skipped levels")
    if len(h["h2"]) < 3:
        issues.append("few H2 sections")
    return "; ".join(issues) or "clear hierarchy"


def a_tone(v, ev) -> dict:
    cp = ev.get("client_page") or {}
    comps = _ok_pages(ev)
    sites = ([cp] if cp.get("status") == 200 else []) + comps
    rows = []
    for p in sites:
        t = p["tone"]
        rows.append({"Site": _host(p["url"]), "Adjectives (measured)": ", ".join(_tone_words(t)), "Avg sentence (words)": t["avg_sentence_words"],
                     "Long sentences %": t["long_sentences_pct"], "Passive %": t["passive_pct"], "Flesch reading ease": t["flesch_reading_ease"],
                     "'You' per 100 words": t["you_per_100w"], "'We' per 100 words": t["we_per_100w"], "Clichés": ", ".join(t["cliches"]) or "none",
                     "CTAs": ", ".join(p["ctas"][:4]), "Sample": (p["sample_paragraphs"][0][:240] if p["sample_paragraphs"] else "")})
    guide = []
    if cp.get("status") == 200:
        t = cp["tone"]
        ls = cp["long_sentences"][0] if cp["long_sentences"] else None
        guide = [
            {"Element": "Voice", "Recommendation": f"Aim for: {v.get('tone') or 'knowledgeable but approachable guide'} — always on the customer's side.", "Example (before → after)": "N/A (defining statement)"},
            {"Element": "Vocabulary", "Recommendation": ("Remove clichés: " + ", ".join(t["cliches"]) + ". ") if t["cliches"] else "Keep plain, specific words; avoid jargon.",
             "Example (before → after)": _cliche_example(cp)},
            {"Element": "Sentence length & structure", "Recommendation": f"Average {t['avg_sentence_words']} words, {t['long_sentences_pct']}% over 25 words, {t['passive_pct']}% passive. "
             "Target 14–18 words average, mostly active voice.", "Example (before → after)": f"Before: “{ls[:200]}” → After: split into two sentences, lead with the action." if ls else "—"},
            {"Element": "Customer focus", "Recommendation": ("Talk to the reader more: 'you' appears less than 'we'." if t["you_per_100w"] < t["we_per_100w"] else "Good reader focus ('you' ≥ 'we')."),
             "Example (before → after)": "Before: “We provide…” → After: “You get…”"},
            {"Element": "Emotional resonance", "Recommendation": "Acknowledge the reader's problem, then give a calm, confident path forward.", "Example (before → after)": "Before: fear-led headline → After: outcome-led headline"},
            {"Element": "CTA integration", "Recommendation": "CTAs as a helpful next step with collaborative verbs.",
             "Example (before → after)": f"Before: “{cp['ctas'][0]}” → After: a specific outcome, e.g. “Get my free plan”" if cp["ctas"] else "Add a primary CTA"},
        ]
    return {"tables": {"Tone audit": rows, "Style guide": guide}, "notes": _need_pages(ev, competitors=False)}


def _tone_words(t):
    w = []
    f = t["flesch_reading_ease"] or 0
    w.append("easy to read" if f >= 60 else ("moderately complex" if f >= 40 else "complex / formal"))
    w.append("reader-focused" if t["you_per_100w"] >= t["we_per_100w"] else "company-focused")
    if (t["passive_pct"] or 0) > 20:
        w.append("impersonal (passive)")
    if t["exclamations"] > 3:
        w.append("enthusiastic")
    if t["cliches"]:
        w.append("corporate")
    if t["questions"] > 3:
        w.append("conversational")
    return w


def _cliche_example(p):
    for s in p["sample_paragraphs"]:
        for c in p["tone"]["cliches"]:
            if c in s.lower():
                i = s.lower().index(c)
                return f"Before: “…{s[max(0, i - 60):i + 60]}…” → After: say what the reader actually gets instead of “{c}”."
    return "—"


def _ngrams(text, n):
    toks = [t for t in re.findall(r"[a-z][a-z0-9\-]+", text.lower())]
    out = []
    for i in range(len(toks) - n + 1):
        g = toks[i:i + n]
        if g[0] in STOP or g[-1] in STOP or any(len(x) < 3 for x in (g[0], g[-1])):
            continue
        if n == 1 and (g[0] in STOP or len(g[0]) < 4):
            continue
        out.append(" ".join(g))
    return out


def a_domain_terms(v, ev) -> dict:
    srcs = {}
    cp = ev.get("client_page") or {}
    if cp.get("status") == 200:
        srcs["our site"] = cp["text"] + " " + " ".join(sum(cp["headings"].values(), []))
    for c in _ok_pages(ev):
        srcs[_host(c["url"])] = c["text"] + " " + " ".join(sum(c["headings"].values(), []))
    q = ev.get("gsc_queries") or []
    if q:
        srcs["search console"] = " . ".join(x["query"] for x in q)
    brand_terms = {b.lower() for b in [v.get("brand")] + [_host(u).split(".")[0] for u in v.get("competitor_urls", [])] if b}
    freq, docs = Counter(), defaultdict(set)
    for name, txt in srcs.items():
        for n in (1, 2, 3):
            for g in _ngrams(txt, n):
                freq[g] += 1
                docs[g].add(name)
    scored = []
    for g, f in freq.items():
        d = len(docs[g])
        if f < 2 and d < 2:
            continue
        if any(b and b in g for b in brand_terms):
            continue
        score = f * (1 + d) * (1.4 if " " in g else 1)
        scored.append((score, g, f, d))
    scored.sort(reverse=True)
    picked, seen = [], set()
    for score, g, f, d in scored:
        if g in seen:
            continue
        seen.add(g)
        picked.append({"Term": g, "Type": ["word", "phrase", "phrase"][min(2, g.count(" "))] if " " in g else "word", "Words": g.count(" ") + 1,
                       "Frequency": f, "Sources": ", ".join(sorted(docs[g])), "Source count": len(docs[g]),
                       "Used on our site": _yn("our site" in docs[g]) if "our site" in srcs else "—",
                       "Intent": kw.intent(g), "Score": round(score, 1)})
        if len(picked) >= int(v.get("limit") or 150):
            break
    notes = [] if srcs else ["No text sources: add the website/competitor URLs or import Search Console data."]
    if picked and len(picked) < 100:
        notes.append(f"Only {len(picked)} recurring terms found in the available text — add more competitor URLs for 100+.")
    return {"tables": {"Domain terms": picked}, "notes": notes}


MODIFIERS = {
    "question": ["how to {k}", "what is {k}", "why {k}", "is {k} worth it", "how much does {k} cost", "how long does {k} take"],
    "comparison": ["{k} vs alternatives", "best {k}", "{k} reviews", "top {k} {year}", "{k} alternatives"],
    "transactional": ["{k} price", "{k} cost", "{k} quote", "hire {k}", "book {k}", "{k} near me"],
    "local": ["{k} in {location}", "{k} {location}", "best {k} in {location}"],
    "audience": ["{k} for {audience}", "{k} for beginners", "{k} for small business"],
    "problem": ["{k} problems", "{k} mistakes to avoid", "{k} checklist", "{k} guide"],
}
FUNNEL = {"informational": "TOFU", "navigational": "BOFU", "commercial": "MOFU", "transactional": "BOFU", "local": "BOFU"}


def _clusters(items: list[tuple[str, int]], min_size=2) -> list[dict]:
    cl = kw.cluster([{"keyword": k, "impressions": i or 0, "intent": kw.intent(k)} for k, i in items], min_size=min_size)
    return [{"name": c["head"], "pillar": c["pillar"], "impressions": c["impressions"], "keywords": [x["keyword"] for x in c["keywords"]]} for c in cl]


def a_keywords(v, ev) -> dict:
    q = ev.get("gsc_queries") or []
    brand_terms = [b.lower() for b in [v.get("brand")] if b] + [_host(u).split(".")[0].lower() for u in v.get("competitor_urls", [])]
    rows, seen = [], set()

    def add(k, source, mod="—", data=None):
        k = re.sub(r"\s+", " ", k.strip().lower())
        if not k or k in seen:
            return
        seen.add(k)
        it = kw.intent(k, v.get("brand"))
        branded = any(b and b in k for b in brand_terms)
        chan = "AEO" if re.match(r"^(how|what|why|when|which|who|is|can|does|should)\b", k) else ("GEO" if len(k.split()) >= 6 else "SEO")
        rows.append({"Keyword": k, "Source": source, "Funnel": FUNNEL.get(it, "MOFU"), "Intent": it, "Modifier": mod,
                     "Branded": "Yes" if branded else "No", "Channel": chan, "Clicks": (data or {}).get("clicks"),
                     "Impressions": (data or {}).get("impressions"), "Position": (data or {}).get("position"),
                     "Validated": "Search Console" if data else "suggestion (not volume-validated)"})
    for x in q:
        add(x["query"], "Search Console", data=x)
    seeds = v.get("primary_keywords") or ([v["primary_query"]] if v.get("primary_query") else [])
    seeds += (v.get("secondary_keywords") or [])[:5]
    for s in seeds:
        add(s, "seed")
        for mod, pats in MODIFIERS.items():
            for p in pats:
                if "{location}" in p and not v.get("location"):
                    continue
                if "{audience}" in p and not v.get("audience"):
                    continue
                add(p.format(k=s, location=v.get("location") or "", audience=v.get("audience") or "", year=v.get("year")), "seed expansion", mod)
    if v.get("exclude_branded", True):
        rows_nb = [r for r in rows if r["Branded"] == "No"]
    else:
        rows_nb = rows
    cl = _clusters([(r["Keyword"], r["Impressions"]) for r in rows_nb])
    cmap = {k: c["name"] for c in cl for k in c.get("keywords", [])}
    for r in rows_nb:
        r["Cluster"] = cmap.get(r["Keyword"], "unclustered")
    clusters = [{"Cluster": c["name"], "Keywords": len(c.get("keywords", [])), "Impressions": c.get("impressions"),
                 "Pillar keyword": c.get("pillar"), "Funnel mix": ", ".join(f"{k}:{n}" for k, n in Counter(FUNNEL.get(kw.intent(x), "MOFU") for x in c.get("keywords", [])).most_common())}
                for c in cl]
    notes = []
    if not q:
        notes.append("No Search Console data for this project — keywords are seed expansions only (not volume-validated). Connect GSC or import a CSV.")
    if not seeds and not q:
        notes.append("Add primary keywords to the project profile or pass a seed keyword.")
    return {"tables": {"Keywords": rows_nb, "Clusters": clusters,
                       "Branded (filtered out)": [r for r in rows if r["Branded"] == "Yes"]}, "notes": notes}


PLAN_COLUMNS = ["A: #", "B: Publish date", "C: Week", "D: Funnel stage", "E: Content type", "F: Working title", "G: Primary keyword",
                "H: Secondary keywords", "I: Search intent", "J: Persona / pain point", "K: Angle & outline", "L: URL slug",
                "M: Internal links to", "N: CTA", "O: KPI / status"]
TYPES = {"TOFU": ["How-to guide", "Explainer", "Checklist", "Listicle", "Glossary entry", "Infographic post"],
         "MOFU": ["Comparison", "Best-of list", "Case study", "Buyer's guide", "Webinar recap", "Template / tool"],
         "BOFU": ["Service / product page", "Pricing explainer", "Local landing page", "Testimonial story", "Demo / consultation page"]}
CTA = {"TOFU": "Subscribe / download checklist", "MOFU": "See case study / compare plans", "BOFU": "Book a consultation / get a quote"}
KPI = {"TOFU": "Impressions + rankings (top 20 in 30 days)", "MOFU": "Clicks + assisted conversions", "BOFU": "Conversions / leads"}


def a_content_plan(v, ev, kw_tables=None) -> dict:
    kws = (kw_tables or a_keywords(v, ev))["tables"]["Keywords"]
    split = {"TOFU": int(v.get("tofu") or 40), "MOFU": int(v.get("mofu") or 30), "BOFU": int(v.get("bofu") or 20)}
    buckets = {s: [r for r in kws if r["Funnel"] == s] for s in split}
    for s in buckets:  # real (validated) keywords first, by impressions
        buckets[s].sort(key=lambda r: (r["Validated"] != "Search Console", -(r["Impressions"] or 0)))
    start = v.get("start_date")
    start = date.fromisoformat(start) if start else date.today() + timedelta(days=(7 - date.today().weekday()) % 7 or 7)
    persona = v.get("persona") or v.get("audience") or "target customer"
    pains = v.get("pain_points") or []
    rows, n = [], 0
    templated = 0
    pillar_by_cluster = {}
    for stage in ("TOFU", "MOFU", "BOFU"):
        pool = buckets[stage]
        for i in range(split[stage]):
            n += 1
            if i < len(pool):
                r = pool[i]
                k, cl, it = r["Keyword"], r.get("Cluster") or "", r["Intent"]
            else:  # not enough keywords for this stage: derive from seeds, labelled
                seed = (v.get("primary_keywords") or [v.get("primary_query") or v.get("product") or "your service"])[i % max(1, len(v.get("primary_keywords") or [1]))]
                pat = {"TOFU": ["what is {k}", "how to choose {k}", "{k} checklist", "{k} mistakes to avoid", "{k} explained for beginners", "{k} myths"],
                       "MOFU": ["best {k} options compared", "{k} case study", "{k} pricing compared", "how to evaluate {k} providers"],
                       "BOFU": ["{k} quote", "{k} in {loc}", "book {k}", "{k} packages"]}[stage][i % (6 if stage == "TOFU" else 4)]
                k = pat.format(k=seed, loc=v.get("location") or "your area")
                cl, it, templated = "seed", {"TOFU": "informational", "MOFU": "commercial", "BOFU": "transactional"}[stage], templated + 1
            ctype = TYPES[stage][i % len(TYPES[stage])]
            title = kw.title_for(k, it, v.get("year") or YEAR)
            sec = [x["Keyword"] for x in kws if x.get("Cluster") == cl and x["Keyword"] != k][:3] if cl and cl != "seed" else []
            slug = "/" + ("blog/" if stage != "BOFU" else "") + re.sub(r"[^a-z0-9]+", "-", k.lower()).strip("-")[:70]
            pillar = pillar_by_cluster.setdefault(cl, slug)
            pain = pains[n % len(pains)] if pains else "—"
            rows.append({PLAN_COLUMNS[0]: n, PLAN_COLUMNS[1]: (start + timedelta(days=n - 1)).isoformat(), PLAN_COLUMNS[2]: (n - 1) // 7 + 1,
                         PLAN_COLUMNS[3]: stage, PLAN_COLUMNS[4]: ctype, PLAN_COLUMNS[5]: title, PLAN_COLUMNS[6]: k,
                         PLAN_COLUMNS[7]: ", ".join(sec), PLAN_COLUMNS[8]: it, PLAN_COLUMNS[9]: f"{persona} — {pain}",
                         PLAN_COLUMNS[10]: _angle(stage, ctype, k), PLAN_COLUMNS[11]: slug,
                         PLAN_COLUMNS[12]: pillar if pillar != slug else "(pillar)", PLAN_COLUMNS[13]: CTA[stage], PLAN_COLUMNS[14]: KPI[stage] + " · Planned"})
    notes = []
    if templated:
        notes.append(f"{templated} of {n} rows are template-derived from seed keywords because there was not enough Search Console data for every stage "
                     "(marked by Cluster = seed in the keyword table). Connect an AI provider to rewrite them, or import more keyword data.")
    return {"tables": {"90-day plan": rows}, "notes": notes, "columns": PLAN_COLUMNS}


def _angle(stage, ctype, k):
    return {"TOFU": f"Answer “{k}” in the first 60 words; H2s for definition, steps, examples, mistakes, FAQ.",
            "MOFU": f"{ctype}: criteria table, pros/cons, proof (numbers, screenshots), who it's for.",
            "BOFU": f"{ctype}: outcome-led H1, proof, process, pricing guidance, objections FAQ, strong CTA."}[stage]


def a_seo_llm_gap(v, ev) -> dict:
    cp = ev.get("client_page") or {}
    if cp.get("status") != 200:
        return {"tables": {}, "notes": _need_pages(ev, competitors=False)}
    comps = _ok_pages(ev)
    kws = v.get("primary_keywords") or ([v["primary_query"]] if v.get("primary_query") else [])
    ts = cp["trust_signals"]
    checks = [
        ("Foundational SEO", "Title 30–60 chars with keyword", 30 <= cp["title_length"] <= 60 and (_kw_in(cp["title"], kws) is not False)),
        ("Foundational SEO", "Meta description 70–160 chars", 70 <= cp["meta_length"] <= 160),
        ("Foundational SEO", "Exactly one H1", cp["heading_counts"]["h1"] == 1),
        ("Foundational SEO", "Logical heading hierarchy", cp["heading_level_skips"] == 0),
        ("Foundational SEO", "Indexable", not cp["noindex"]),
        ("Foundational SEO", "Mobile viewport", bool(cp.get("viewport"))),
        ("Foundational SEO", "Internal links (≥10)", (cp.get("internal_links") or 0) >= 10),
        ("Foundational SEO", "Images have alt text", cp["images_missing_alt"] == 0),
        ("NLP & entity clarity", "Defines key concepts ('X is…')", cp["definitions"] >= 2),
        ("NLP & entity clarity", "Question-style headings", len(cp["faq_questions"]) >= 3),
        ("NLP & entity clarity", "Lists/tables for extractable answers", cp["lists"] + cp["tables"] >= 3),
        ("NLP & entity clarity", "Structured data present", bool(cp["schema_types"])),
        ("NLP & entity clarity", "FAQ schema", cp["faq_schema"]),
        ("NLP & entity clarity", "Readable (Flesch ≥ 50)", (cp["tone"]["flesch_reading_ease"] or 0) >= 50),
        ("E-E-A-T: Experience", "First-hand proof (case studies / results)", ts["case_studies"] or bool(re.search(r"\b(we|our team) (tested|helped|built|completed)\b", cp["text"].lower()))),
        ("E-E-A-T: Expertise", "Author / credentials shown", ts["author"] or ts["certifications"]),
        ("E-E-A-T: Authoritativeness", "Awards, ratings or external citations", ts["awards"] or ts["ratings"] or (cp.get("external_links") or 0) >= 3),
        ("E-E-A-T: Trust", "Contact details (phone/address)", ts["phone"] or ts["address"]),
        ("E-E-A-T: Trust", "Privacy/security signals", ts["privacy_security"]),
        ("E-E-A-T: Trust", "Last-updated date", bool(cp["date_modified"])),
    ]
    rows = [{"Area": a, "Check": c, "Result": "Pass" if ok else "Gap",
             "Competitors passing": _comp_pass(comps, c) if comps else "—"} for a, c, ok in checks]
    ymyl = cp["ymyl"]
    recs = []
    pri = {"Foundational SEO": "High", "E-E-A-T: Trust": "High" if ymyl else "Medium"}
    for r in rows:
        if r["Result"] == "Gap":
            recs.append({"Priority": pri.get(r["Area"], "High" if ymyl and r["Area"].startswith("E-E-A-T") else "Medium"), "Gap": r["Check"], "Area": r["Area"], "Action": _gap_action(r["Check"])})
    recs.sort(key=lambda x: x["Priority"] != "High")
    if comps:
        ours = {kw.stem(t) for t in kw.tokens(cp["text"])}
        common = Counter(kw.stem(t) for c in comps for t in set(kw.tokens(c["text"])))
        missing = [t for t, n in common.most_common(300) if n >= max(2, len(comps) // 2) and t not in ours and len(t) > 3][:30]
        if missing:
            recs.append({"Priority": "Medium", "Gap": "Missing semantic terms", "Area": "NLP & entity clarity", "Action": "Cover: " + ", ".join(missing)})
    return {"tables": {"Checks": rows, "Priority recommendations": recs[:12],
                       "YMYL": [{"YMYL topic": "Yes" if ymyl else "No", "Why": "health/finance/legal/safety vocabulary detected" if ymyl else "no YMYL vocabulary detected"}]},
            "notes": []}


def _comp_pass(comps, check):
    m_ = {"Exactly one H1": lambda p: p["heading_counts"]["h1"] == 1, "FAQ schema": lambda p: p["faq_schema"],
          "Structured data present": lambda p: bool(p["schema_types"]), "Question-style headings": lambda p: len(p["faq_questions"]) >= 3,
          "Author / credentials shown": lambda p: p["trust_signals"]["author"] or p["trust_signals"]["certifications"],
          "Last-updated date": lambda p: bool(p["date_modified"]), "Readable (Flesch ≥ 50)": lambda p: (p["tone"]["flesch_reading_ease"] or 0) >= 50}
    f = m_.get(check)
    return f"{sum(1 for p in comps if f(p))}/{len(comps)}" if f else "—"


def _gap_action(c):
    return {"Title 30–60 chars with keyword": "Rewrite the title (50–60 chars, keyword first, differentiator).",
            "Meta description 70–160 chars": "Rewrite the meta description (benefit, proof, CTA; 140–155 chars).",
            "Exactly one H1": "Use one H1 that states the page's main promise.", "Logical heading hierarchy": "Nest headings H2→H3→H4 without skipping.",
            "Indexable": "Remove noindex.", "Mobile viewport": "Add a responsive viewport meta tag.", "Internal links (≥10)": "Add contextual internal links to/from related pages.",
            "Images have alt text": "Add descriptive alt text.", "Defines key concepts ('X is…')": "Add one-sentence definitions of the key entities near the top.",
            "Question-style headings": "Add question H2/H3s matching real search questions, each answered in 40–60 words.",
            "Lists/tables for extractable answers": "Add steps, checklists and a comparison table answer engines can quote.",
            "Structured data present": "Add JSON-LD (Organization, Service/Product, Article, BreadcrumbList).", "FAQ schema": "Add FAQPage schema for the FAQ block.",
            "Readable (Flesch ≥ 50)": "Shorten sentences and replace jargon.", "First-hand proof (case studies / results)": "Add first-hand results: numbers, photos, client outcomes.",
            "Author / credentials shown": "Add an author byline with credentials and a bio page.", "Awards, ratings or external citations": "Cite authoritative sources; show ratings/awards.",
            "Contact details (phone/address)": "Show phone, address and contact page in the footer and schema.", "Privacy/security signals": "Link privacy policy; show security/compliance badges.",
            "Last-updated date": "Show a last-updated date and dateModified in schema."}.get(c, "Fix this gap.")


def a_semantic(v, ev) -> dict:
    pq = (v.get("primary_query") or "").lower()
    q = ev.get("gsc_queries") or []
    toks = set(kw.stem(t) for t in kw.tokens(pq))
    pasf = [x for x in q if toks & {kw.stem(t) for t in kw.tokens(x["query"])} and x["query"] != pq][:100]
    pasf_rows = [{"Query": x["query"], "Type": _qtype(x["query"]), "Intent": kw.intent(x["query"]), "Impressions": x["impressions"], "Position": x["position"]} for x in pasf]
    templates = Counter()
    for x in pasf:
        t = x["query"]
        for tk in kw.tokens(pq):
            t = re.sub(rf"\b{re.escape(tk)}\w*\b", "{entity}", t)
        if v.get("location"):
            t = t.replace(v["location"].lower(), "{location}")
        if "{entity}" in t:
            templates[re.sub(r"(\{entity\}\s*)+", "{entity} ", t).strip()] += 1
    comps = _ok_pages(ev)
    cp = ev.get("client_page") or {}
    texts = [c["text"] for c in comps] + ([cp["text"]] if cp.get("status") == 200 else [])
    ent = Counter()
    for t in texts:
        for g in set(_ngrams(t, 2) + _ngrams(t, 1)):
            ent[g] += 1
    entities = [{"Entity / concept": e, "Pages mentioning": n, "On our page": _yn(cp.get("status") == 200 and e in cp["text"].lower())}
                for e, n in ent.most_common(60) if n >= max(2, len(texts) // 2)][:40]
    cl = _clusters([(x["query"], x["impressions"]) for x in pasf])
    return {"tables": {"Intent": [{"Primary query": pq, "Primary intent": kw.intent(pq) if pq else "—", "Funnel": FUNNEL.get(kw.intent(pq), "MOFU") if pq else "—"}],
                       "PASF (from Search Console)": pasf_rows,
                       "Query templates (observed)": [{"Template": t, "Occurrences": n} for t, n in templates.most_common(50)],
                       "Entities (from ranking pages)": entities,
                       "Topic clusters": [{"Cluster": c["name"], "Keywords": ", ".join(c.get("keywords", [])[:12])} for c in cl]},
            "notes": ([] if q else ["No Search Console data — PASF and templates need GSC queries or an AI provider."]) +
                     ([] if pq else ["Set a primary query (run input) or primary keywords in the profile."]) +
                     ["Sub-entities, attributes, context layers and perspectives are written by the AI step (connect an AI provider)."]}


def _qtype(q):
    q = q.lower()
    if re.match(r"^(how|what|why|when|where|who|which|is|can|does|should)\b", q):
        return "question"
    if " vs " in q or "versus" in q or "compare" in q:
        return "comparison"
    if "near me" in q or re.search(r"\bin [a-z]+$", q):
        return "local"
    if re.search(r"price|cost|cheap|quote", q):
        return "product/service"
    if re.search(r"problem|issue|not working|fix|error", q):
        return "problem-based"
    return "long-tail" if len(q.split()) >= 4 else "exact-match"


GEO_COLUMNS = ["Query ID", "Search Query", "Search Intent", "Query Type (Traditional / AI-Conversational)", "Related Service / Category",
               "Industry Context", "User Situation / Trigger", "Urgency Level (Low / Medium / High)",
               "Likely AI Response Style (Summary / Comparison / Advice)", "Zero-Click Potential (Low / Medium / High)"]
GEO_T = [  # (pain model, mode, intent, type, urgency, style, zero-click, situation, template)
    ("Trust deficit", "Trust & verification", "Trust & Verification", "AI-Conversational", "Medium", "Advice", "Medium", "Choosing a provider",
     "I'm in {area} and about to hire someone for {service}.\nWhat questions should I ask before I trust them, and how do I check their credentials?"),
    ("Trust deficit", "Trust & verification", "Trust & Verification", "AI-Conversational", "Medium", "Advice", "Medium", "Got a quote",
     "I got a quote for {service} that seems vague.\nWhat costs or risks are usually not explained upfront?"),
    ("Trust deficit", "Trust & verification", "Trust & Verification", "Traditional", "Low", "Summary", "High", "Researching",
     "how do i know if a {service} provider is legit"),
    ("Trust deficit", "Transactional but cautious", "Transactional", "AI-Conversational", "Medium", "Advice", "Low", "Ready to book",
     "I need {service} soon but I've been burned before.\nWhat red flags should make me walk away?"),
    ("Trust deficit", "Trust & verification", "Trust & Verification", "AI-Conversational", "Low", "Advice", "Medium", "Reading reviews",
     "Can I trust online reviews for {service} companies, or how do I spot fake ones?"),
    ("Trust deficit", "Trust & verification", "Commercial Investigation", "AI-Conversational", "Low", "Comparison", "Medium", "Comparing quotes",
     "I have three quotes for {service} in {area} and they're very different.\nHow do I tell which one is honest?"),
    ("Uncertainty & trade-offs", "Comparison & alternatives", "Comparison", "AI-Conversational", "Medium", "Comparison", "Medium", "Tight budget",
     "My budget is around {budget}.\nIs the cheaper {service} option actually riskier in the long run?"),
    ("Uncertainty & trade-offs", "Comparison & alternatives", "Comparison", "AI-Conversational", "Low", "Comparison", "High", "Considering DIY",
     "Should I try to handle this myself or pay for {service}?\nWhat could go wrong if I do it myself?"),
    ("Uncertainty & trade-offs", "Problem-first", "Exploratory / Need-Based", "AI-Conversational", "Medium", "Advice", "Medium", "Putting it off",
     "Does waiting make it worse if I don't get {service} right away?"),
    ("Uncertainty & trade-offs", "Educational & preventative", "Educational", "Traditional", "Low", "Summary", "High", "Researching",
     "is {service} worth the money"),
    ("Uncertainty & trade-offs", "Comparison & alternatives", "Comparison", "AI-Conversational", "Low", "Comparison", "Medium", "Weighing options",
     "Can you compare the main options for {service} and tell me which suits someone who wants it done once and properly?"),
    ("Uncertainty & trade-offs", "AI overview / zero-click", "Informational", "Traditional", "Low", "Summary", "High", "First look",
     "what does {service} actually involve"),
    ("Uncertainty & trade-offs", "Problem-first", "Exploratory / Need-Based", "AI-Conversational", "Medium", "Advice", "Medium", "Overwhelmed",
     "I don't know what I don't know about {service}.\nWhat are the things people usually regret not asking?"),
    ("Dependent & family safety", "Risk & side effects", "Informational", "AI-Conversational", "High", "Advice", "Medium", "Has young kids",
     "I have a toddler at home.\nIs {service} safe around young children, and what precautions should we take?"),
    ("Dependent & family safety", "Risk & side effects", "Informational", "AI-Conversational", "Medium", "Advice", "Medium", "Elderly parent",
     "My elderly mother lives with us in {area}.\nAre there any risks from {service} I should know about for older people?"),
    ("Dependent & family safety", "Risk & side effects", "Informational", "Traditional", "Medium", "Summary", "High", "Pets at home",
     "is {service} safe for pets"),
    ("Dependent & family safety", "Risk & side effects", "Informational", "AI-Conversational", "High", "Advice", "Medium", "Pregnancy",
     "My partner is pregnant.\nShould we delay {service} or are there safer ways to do it?"),
    ("Dependent & family safety", "Risk & side effects", "Informational", "AI-Conversational", "Medium", "Advice", "Medium", "Health condition",
     "Someone in my home has a weak immune system.\nWhat should I ask the {service} provider before they start?"),
    ("Dependent & family safety", "Educational & preventative", "Educational", "Traditional", "Low", "Summary", "High", "Planning",
     "how to keep kids safe during {service}"),
    ("Values & ethics", "Trust & verification", "Exploratory / Need-Based", "AI-Conversational", "Low", "Advice", "Medium", "Eco-conscious",
     "Are there more environmentally friendly ways to do {service}?\nI don't want to cause long-term harm."),
    ("Values & ethics", "Comparison & alternatives", "Comparison", "Traditional", "Low", "Comparison", "Medium", "Ethical buyer",
     "is there a more humane or gentler alternative to {service}"),
    ("Values & ethics", "Educational & preventative", "Educational", "AI-Conversational", "Low", "Summary", "Medium", "Worried about materials",
     "What materials or chemicals are typically used in {service}, and are any of them a concern?"),
    ("Values & ethics", "Trust & verification", "Trust & Verification", "AI-Conversational", "Low", "Advice", "Low", "Values-driven",
     "I want a {service} provider in {area} who works sustainably.\nHow do I check if their claims are real?"),
    ("Post-service anxiety", "Post-service expectations", "Post-Purchase", "AI-Conversational", "Medium", "Advice", "High", "Just finished",
     "I just had {service} done yesterday.\nWhat's normal to notice in the next few days, and what isn't?"),
    ("Post-service anxiety", "Post-service expectations", "Post-Purchase", "Traditional", "Medium", "Summary", "High", "After the job",
     "how long after {service} is it safe"),
    ("Post-service anxiety", "Post-service expectations", "Post-Purchase", "AI-Conversational", "Medium", "Advice", "Medium", "Doubting results",
     "How can I tell if the {service} actually worked?"),
    ("Post-service anxiety", "Post-service expectations", "Post-Purchase", "AI-Conversational", "High", "Advice", "Medium", "Symptoms appeared",
     "It's been a week since {service} and something doesn't seem right.\nWhen should I call them back or get a second opinion?"),
    ("Post-service anxiety", "Post-service expectations", "Post-Purchase", "AI-Conversational", "Low", "Advice", "Medium", "Maintenance",
     "What should I do after {service} so the results last as long as possible?"),
    ("Emergency & panic", "Symptom & trigger-based", "Emergency / Urgent", "AI-Conversational", "High", "Advice", "High", "Sudden problem",
     "Something just went wrong and I think I need {service} urgently.\nWhat should I do right now before help arrives?"),
    ("Emergency & panic", "Symptom & trigger-based", "Emergency / Urgent", "Traditional", "High", "Summary", "High", "Emergency",
     "is this an emergency or can {service} wait until monday"),
    ("Emergency & panic", "Symptom & trigger-based", "Emergency / Urgent", "AI-Conversational", "High", "Advice", "Medium", "After hours",
     "It's late at night in {area} and I need {service}.\nWho should I contact first, and what will it likely cost?"),
    ("Emergency & panic", "Symptom & trigger-based", "Emergency / Urgent", "AI-Conversational", "High", "Advice", "High", "Panic",
     "I'm panicking — what are the first three things I should do while I wait for {service}?"),
    ("Prevention & long-term", "Educational & preventative", "Educational", "AI-Conversational", "Low", "Advice", "High", "Recurring issue",
     "This keeps happening every year.\nHow do I prevent needing {service} again?"),
    ("Prevention & long-term", "Educational & preventative", "Educational", "Traditional", "Low", "Summary", "High", "Planning ahead",
     "how often should you get {service}"),
    ("Prevention & long-term", "Educational & preventative", "Educational", "AI-Conversational", "Low", "Advice", "Medium", "Long-term thinking",
     "Is ongoing maintenance safer and cheaper than waiting for {service} when things break?"),
    ("Prevention & long-term", "AI overview / zero-click", "Informational", "AI-Conversational", "Low", "Summary", "High", "Want a checklist",
     "Can you give me a simple checklist to reduce the chances I'll need {service}?"),
    ("Regulation & compliance", "Trust & verification", "Trust & Verification", "AI-Conversational", "Medium", "Summary", "Medium", "Checking rules",
     "What regulations apply to {service} in {area}, and which certifications actually matter?"),
    ("Regulation & compliance", "Trust & verification", "Trust & Verification", "Traditional", "Low", "Summary", "High", "Legal question",
     "is {service} legally required"),
    ("Regulation & compliance", "Transactional but cautious", "Transactional", "AI-Conversational", "Medium", "Advice", "Low", "Liability worry",
     "If something goes wrong during {service}, who is liable — me or the company?\nWhat should be in the contract?"),
    ("Regulation & compliance", "Trust & verification", "Trust & Verification", "AI-Conversational", "Low", "Advice", "Medium", "Insurance",
     "Do I need to check a {service} company's insurance, and how do I verify it?"),
]
GEO_STARTERS = ["", "Quick question: ", "Honestly, ", "Be straight with me — ", "For context, "]


def a_geo_dataset(v, ev) -> dict:
    services = (v.get("primary_keywords") or []) + (v.get("secondary_keywords") or [])
    services = services or [v.get("product") or v.get("primary_query") or ""]
    services = [s for s in services if s][:4]
    if not services:
        return {"tables": {}, "notes": ["Set the service (product / primary keywords) in the profile or run inputs."]}
    area = v.get("location") or "my area"
    budget = v.get("budget_note") or "a limited budget"
    q = ev.get("gsc_queries") or []
    rows, seen, n = [], set(), 0
    target = int(v.get("count") or 120)
    i = 0
    while len(rows) < target and i < len(GEO_T) * len(services) * len(GEO_STARTERS):
        tpl = GEO_T[i % len(GEO_T)]
        svc = services[(i // len(GEO_T)) % len(services)]
        start = GEO_STARTERS[(i // (len(GEO_T) * len(services))) % len(GEO_STARTERS)]
        i += 1
        text = tpl[8].format(service=svc, area=area, budget=budget)
        if start:
            text = start + text[0].lower() + text[1:] if not text.startswith("I") else start + text
        if text in seen:
            continue
        seen.add(text)
        n += 1
        rows.append({GEO_COLUMNS[0]: f"Q{n:03d}", GEO_COLUMNS[1]: text, GEO_COLUMNS[2]: tpl[2], GEO_COLUMNS[3]: tpl[3], GEO_COLUMNS[4]: svc,
                     GEO_COLUMNS[5]: v.get("industry") or "—", GEO_COLUMNS[6]: f"{tpl[0]}: {tpl[7]}", GEO_COLUMNS[7]: tpl[4],
                     GEO_COLUMNS[8]: tpl[5], GEO_COLUMNS[9]: tpl[6]})
    real = [x for x in q if len(x["query"].split()) >= 5 or re.match(r"^(how|what|why|is|can|should|do|which)\b", x["query"])][:30]
    return {"tables": {"GEO query dataset": rows,
                       "Real conversational queries (Search Console)": [{"Query": x["query"], "Impressions": x["impressions"], "Position": x["position"]} for x in real]},
            "notes": ["Rows are generated from the 8 pain-model templates (all 9 query modes covered) — labelled template-based. "
                      "With an AI provider connected the full GEO prompt runs and writes naturally varied queries."], "columns": GEO_COLUMNS}


EEAT_COLUMNS = ["Title ID", "Main Title (≤70 characters)", "Subtitle (scope + benefit)", "Intent Type", "E-E-A-T Focus", "Target Audience Persona",
                "Primary Pain Point", "Verifiable Metric / Result", "Core Keyword", "2025–2026 Trend Cue", "Experience Proof Element",
                "Expertise Demonstration", "Authority Signal", "Trustworthiness Indicator", "SEO-Optimized Variant", "Social Media Hook Variant",
                "Email Subject Line Variant"]
EEAT_F = [
    ("Experience", "Problem-Solving", "How We {action} for {industry}: A {year} Implementation Study"),
    ("Experience", "Educational", "Our {months}-Month Hands-On Journey Solving {pain}"),
    ("Experience", "Problem-Solving", "We Documented Everything: A {industry} {topic} Transformation"),
    ("Expertise", "Educational", "Advanced {topic} Implementation for {industry}"),
    ("Expertise", "Strategic", "The {solution} Framework Explained: A {year} Deep Dive"),
    ("Expertise", "Provocative", "Why Old {topic} Tactics Failed and Our {solution} Worked"),
    ("Authoritativeness", "Strategic", "Setting the {year} {industry} Standard for {topic}"),
    ("Authoritativeness", "Educational", "The Definitive {year} Guide to {topic} for {audience}"),
    ("Authoritativeness", "Strategic", "How {industry} Leaders Build Authority with {topic}"),
    ("Trustworthiness", "Commercial", "{metric} in {months} Months: A Data-Verified {topic} Case Study"),
    ("Trustworthiness", "Commercial", "From {pain} to {metric}: A {months}-Month {topic} Story"),
    ("Trustworthiness", "Commercial", "Cutting {pain} with {solution}: Verified {year} Results"),
    ("Experience", "Problem-Solving", "The {pain} Crisis: How We Built a Verified Fix"),
    ("Expertise", "Problem-Solving", "Navigating {pain}: A Step-by-Step {topic} Case Study"),
    ("Authoritativeness", "Provocative", "Solving {industry}'s {pain} Problem in {year}"),
]
TREND = ["AI-assisted, human-verified", "Zero-click / AI Overview ready", "Data privacy first", "First-party data", "ESG & sustainability",
         "Answer-engine visibility", "Security by design", "Automation at scale"]


def a_eeat_titles(v, ev) -> dict:
    topic = v.get("topic") or v.get("primary_query") or (v.get("primary_keywords") or [None])[0] or v.get("product")
    if not topic:
        return {"tables": {}, "notes": ["Set a core topic (run input) or primary keywords in the profile."]}
    d = {"topic": topic.title() if topic.islower() else topic, "industry": v.get("industry") or "B2B", "year": v.get("year") or YEAR,
         "pain": (v.get("pain_points") or ["Rising Costs"])[0].title(), "solution": v.get("solution") or v.get("product") or "Our Method",
         "metric": v.get("metric") or (v.get("proof_points") or "+40% ROI").split(",")[0].strip(), "months": str(v.get("months") or 6),
         "audience": v.get("audience") or "Decision-Makers", "action": v.get("action") or f"Fixed {topic}"}
    rows = []
    count = int(v.get("count") or 15)
    for i, (focus, intent, f) in enumerate(EEAT_F[:count]):
        t = f.format(**d)
        if len(t) > 70:
            t = t[:67].rsplit(" ", 1)[0] + "…"
        rows.append({EEAT_COLUMNS[0]: f"T{i + 1:02d}", EEAT_COLUMNS[1]: t, EEAT_COLUMNS[2]: f"How {d['audience']} in {d['industry']} can tackle {d['pain'].lower()} with {d['solution']} — scope, steps and verified results",
                     EEAT_COLUMNS[3]: intent, EEAT_COLUMNS[4]: focus, EEAT_COLUMNS[5]: d["audience"], EEAT_COLUMNS[6]: d["pain"], EEAT_COLUMNS[7]: d["metric"],
                     EEAT_COLUMNS[8]: topic, EEAT_COLUMNS[9]: TREND[i % len(TREND)],
                     EEAT_COLUMNS[10]: "Before/after screenshots and timeline from our own project", EEAT_COLUMNS[11]: f"Named methodology: {d['solution']}",
                     EEAT_COLUMNS[12]: "Benchmarks vs industry averages; expert quotes", EEAT_COLUMNS[13]: "Data sources and method disclosed; client-approved",
                     EEAT_COLUMNS[14]: f"{topic}: {t.split(':')[0]} ({d['year']})"[:70], EEAT_COLUMNS[15]: f"{d['metric']} in {d['months']} months. Here's exactly how 👇",
                     EEAT_COLUMNS[16]: f"[Case study] {d['metric']} — {topic}"[:60]})
    return {"tables": {"E-E-A-T titles": rows}, "columns": EEAT_COLUMNS,
            "notes": ["Titles are formula-based from your inputs (metric, pain point, solution) — check every metric is verifiable before publishing. "
                      "With an AI provider connected the full generator prompt runs for more varied angles."]}


def a_icp(v, ev) -> dict:
    q = ev.get("gsc_queries") or []
    pains = [x for x in q if re.search(r"problem|issue|not working|why|cost|expensive|slow|risk|safe|mistake|fix|help", x["query"])][:40]
    dist = Counter(kw.intent(x["query"]) for x in q)
    rows = [{"Signal": "Search intent mix", "Evidence": ", ".join(f"{k} {n}" for k, n in dist.most_common()) or "no data"},
            {"Signal": "Profile audience", "Evidence": v.get("audience") or "not set"},
            {"Signal": "Profile persona", "Evidence": v.get("persona") or "not set"},
            {"Signal": "Stated pain points", "Evidence": ", ".join(v.get("pain_points") or []) or "not set"},
            {"Signal": "Location", "Evidence": v.get("location") or "not set"}]
    cp = ev.get("client_page") or {}
    if cp.get("status") == 200:
        rows.append({"Signal": "Who our page talks to", "Evidence": f"'you' {cp['tone']['you_per_100w']}/100w vs 'we' {cp['tone']['we_per_100w']}/100w; hero: {cp['hero'][:150]}"})
    return {"tables": {"Audience signals": rows,
                       "Pain-point queries (Search Console)": [{"Query": x["query"], "Impressions": x["impressions"], "Intent": kw.intent(x["query"])} for x in pains]},
            "notes": ["The ICP narrative (demographics, firmographics, jobs-to-be-done, objections, buying triggers, channels) is written by the AI step "
                      "from these signals — connect an AI provider. The signals above are real data."]}


def a_faq(v, ev) -> dict:
    q = [x for x in ev.get("gsc_queries") or [] if re.match(r"^(how|what|why|when|where|who|which|is|are|can|does|do|should)\b", x["query"])][:40]
    cp = ev.get("client_page") or {}
    have = {h.lower() for h in cp.get("faq_questions", [])} if cp.get("status") == 200 else set()
    rows = [{"Question": x["query"].capitalize().rstrip("?") + "?", "Impressions": x["impressions"], "Position": x["position"],
             "Already on page": _yn(any(kw.tokens(x["query"])[0] in h for h in have) if have and kw.tokens(x["query"]) else False)} for x in q]
    schema = {"@context": "https://schema.org", "@type": "FAQPage", "mainEntity": [
        {"@type": "Question", "name": r["Question"], "acceptedAnswer": {"@type": "Answer", "text": "(answer to write — 40–60 words, answer first)"}} for r in rows[:10]]}
    return {"tables": {"FAQ questions": rows}, "schema": schema,
            "notes": [] if q else ["No question queries in Search Console data yet — import/connect GSC."]}


def a_local_pages(v, ev) -> dict:
    locs = _list(v.get("locations") or v.get("location"))
    svcs = (v.get("primary_keywords") or []) + (v.get("secondary_keywords") or [])
    if not locs or not svcs:
        return {"tables": {}, "notes": ["Set locations (comma-separated) and primary/secondary keywords (services)."]}
    q = {x["query"]: x for x in ev.get("gsc_queries") or []}
    rows = []
    for s in svcs[:10]:
        for l in locs[:25]:
            k = f"{s} in {l}".lower()
            d = q.get(k) or q.get(f"{s} {l}".lower())
            rows.append({"Service": s, "Location": l, "Target keyword": k, "URL": "/" + re.sub(r"[^a-z0-9]+", "-", f"{s}-{l}".lower()).strip("-"),
                         "H1": f"{s.title()} in {l}", "Title": f"{s.title()} in {l} | {v.get('brand') or ''}".strip(" |")[:60],
                         "GSC impressions": d["impressions"] if d else None, "Unique local proof to add": "local projects, reviews from the area, map, local phone"})
    return {"tables": {"Local landing pages": rows}, "notes": ["Only publish a location page when you have genuinely local proof (avoid doorway pages)."]}


def a_blog_titles(db, project, v, ev) -> dict:
    ideas = kw.blog_ideas(db, project, limit=int(v.get("limit") or 30))
    rows = [{"Title": i["title"], "Primary keyword": i["primary_keyword"], "Supporting": ", ".join(i.get("supporting_keywords", [])[:5]),
             "Intent": i["intent"], "Demand (impressions)": i.get("demand"), "Format": i.get("format"), "Source": ", ".join(i.get("sources", []))} for i in ideas]
    return {"tables": {"Blog ideas": rows}, "notes": [] if rows else ["No query data yet — connect/import Search Console."]}


# ------------------------------------------------------------------ the skill library
def S(key, name, category, ws, desc, evidence, template, analysis, columns=None, inputs=None, ai_only_parts=None):
    return {"key": key, "name": name, "category": category, "workstream": ws, "description": desc, "evidence": evidence,
            "template": template.strip(), "analysis": analysis, "output_columns": columns or [], "inputs": inputs or [], "builtin": True}


SKILLS = [
    S("meta_gap_analysis", "Meta title & description gap analysis", "On-page", "on_page",
      "Compares your title/meta, USP, tone, content & E-E-A-T, headings, schema, technical and UX/CTA against competitors; SWOT and top 5 actions.",
      ["client_page", "competitor_pages", "gsc_queries"],
      """Act as a senior SEO and conversion copywriter. Using ONLY the measured evidence below, perform a Meta Title & Description gap analysis for {{brand}} ({{url}}) in {{industry}} against the competitor pages.
Target keywords: {{primary_keywords}}. Audience: {{audience}}. USP: {{usp}}. Desired tone: {{tone}}.
Cover: 1) title & meta (length, keyword placement, CTR hooks, proof), 2) USP & differentiation, 3) tone, 4) content & E-E-A-T, 5) heading structure, 6) schema, 7) technical, 8) UX & CTA,
9) a SWOT, 10) the top 5 prioritised actions. Then write 3 alternative titles (≤60 chars) and 3 meta descriptions (≤155 chars) for our page.
Return markdown tables where comparison is involved.
EVIDENCE:
{{EVIDENCE}}""", "meta_gap", inputs=["url", "competitor_urls", "primary_keywords"]),
    S("icp_deep_dive", "ICP deep dive", "Strategy", "strategy", "Ideal customer profile from your profile, search behaviour and page messaging.",
      ["client_page", "gsc_queries"],
      """Act as a B2B/B2C market strategist. Build an Ideal Customer Profile deep dive for {{brand}} — {{product}} in {{industry}} ({{stage}}, {{revenue_model}}), serving {{audience}} in {{location}}.
Include: demographics/firmographics, roles & decision power, goals, jobs-to-be-done, pain points (use the real search queries as evidence), objections, buying triggers, decision criteria,
information sources & channels, messaging angles that resonate, disqualifiers, and 3 named personas. Base claims on the evidence; mark assumptions.
EVIDENCE:
{{EVIDENCE}}""", "icp"),
    S("domain_terms", "Domain terms extractor (100+)", "Keywords", "content", "Extracts 100+ recurring industry terms and phrases from your page, competitors and search queries.",
      ["client_page", "competitor_pages", "gsc_queries"],
      """Act as a domain linguist and SEO researcher for {{industry}}. From the evidence, produce 100+ domain-specific search terms, entities and phrases used in this niche (exclude brand names),
grouped by category (products/services, problems, processes, materials/tools, metrics, roles, regulations, comparisons, local), each with intent and funnel stage, as a markdown table:
| Term | Category | Intent | Funnel | Used on our site |
EVIDENCE:
{{EVIDENCE}}""", "domain_terms"),
    S("competitor_page_analysis", "Competitor page analysis", "Competitors", "content",
      "Layout, quality signals, visuals, FAQ, comparison matrix, gaps and recommendations for competitor pages vs yours.",
      ["client_page", "competitor_pages"],
      """Act as an expert competitive SEO analyst. Analyse the competitor pages vs our page ({{url}}) for {{primary_keywords}}:
1) page layout & section flow, 2) content quality signals, 3) visuals & media, 4) FAQ coverage, 5) a feature comparison matrix (markdown table), 6) content gaps, 7) prioritised recommendations to outrank them.
Use the measured evidence only.
EVIDENCE:
{{EVIDENCE}}""", "competitor_pages"),
    S("advanced_keyword_discovery", "Advanced keyword discovery (SEO + GEO + AEO)", "Keywords", "on_page",
      "Keyword universe with funnel stage, intent, modifiers, branded filter, clusters and channel (SEO/AEO/GEO).",
      ["gsc_queries", "keywords"],
      """Act as a keyword research expert for SEO, AEO (answer engines) and GEO (generative engines). Seed: {{primary_keywords}} {{secondary_keywords}}; product: {{product}}; audience: {{audience}}; location: {{location}}.
Using the real Search Console queries as ground truth, expand into a full keyword universe: head, long-tail, question (AEO), conversational prompts (GEO), comparison, local, problem-based.
Filter out branded terms ({{brand}} and competitors). Return a markdown table:
| Keyword | Funnel (TOFU/MOFU/BOFU) | Intent | Modifier | Channel (SEO/AEO/GEO) | Cluster | Priority |
EVIDENCE:
{{EVIDENCE}}""", "keywords"),
    S("content_plan_90", "90-day content & SEO plan (Excel)", "Content", "content",
      "90 rows (40 TOFU / 30 MOFU / 20 BOFU), columns A–O, from your real keywords and clusters — downloadable as Excel.",
      ["gsc_queries", "keywords"],
      """Act as a content strategist. Create a 90-day content & SEO plan for {{brand}} ({{product}}, {{industry}}) for {{audience}} in {{location}}.
Exactly 90 rows: 40 TOFU, 30 MOFU, 20 BOFU, one per day starting {{start_date}}. Use the keyword and cluster evidence; keep internal links to pillar pages.
Return ONE markdown table with exactly these columns: {{COLUMNS}}
EVIDENCE:
{{EVIDENCE}}""", "content_plan", columns=PLAN_COLUMNS),
    S("seo_llm_gap", "SEO & LLM gap analysis", "AI search", "geo",
      "Foundational SEO, NLP/entity clarity, YMYL and E-E-A-T evidence audit; variants: simple, detailed, comparative, fill-the-gaps rewrite.",
      ["client_page", "competitor_pages"],
      """Act as a senior SEO and AI Optimization Specialist. Variant: {{variant}}.
Analyse {{url}} for gaps between traditional SEO and LLM optimisation: 1) foundational SEO (meta, headings, internal links, mobile, speed), 2) content quality vs competitors,
3) semantic richness & NLP clarity (entities and relationships, conversational answers, missing semantic terms), 4) YMYL status with justification, 5) E-E-A-T evidence (experience, expertise, authoritativeness, trust),
6) the top 3–5 gaps as High/Medium priority with a concrete action each. If the variant is "comparative", benchmark against the competitor page; if "fill_gaps", rewrite the main content to close the gaps.
EVIDENCE:
{{EVIDENCE}}""", "seo_llm_gap", inputs=["variant"]),
    S("usp_heading_gap", "USP & heading-structure gap analysis", "On-page", "content",
      "Core USP & messaging, content flow & tone, and H1–H6 structure for your site vs up to 10 competitors.",
      ["client_page", "competitor_pages"],
      """Act as an expert content and SEO strategist. For our site and each competitor (up to 10) analyse: A) core USP & messaging (primary promise, shown immediately?, differentiation),
B) content flow & tone (whole-site tone, journey from awareness to action, homepage section-by-section purpose), C) H1–H6 structure (exact H1 and its strength, all H2s and the story they tell, H3–H6 support).
Then fill the gap tables: 1) Core USP & messaging, 2) Heading tag structure, with "Gap & recommendation" columns.
EVIDENCE:
{{EVIDENCE}}""", "usp_heading"),
    S("tone_style_guide", "Tone audit & style guide", "Content", "content",
      "Measured tone of your copy vs competitors (readability, sentence length, passive voice, you/we, clichés) and a style guide with before→after examples.",
      ["client_page", "competitor_pages"],
      """Conduct a tone-of-voice analysis for {{brand}} ({{url}}). Target audience: {{audience}}. Brand goal/promise: {{brand_promise}}.
1) Tone audit: 3–5 adjectives each supported by an example from the text; alignment with audience and brand goal; 2–3 refinements.
2) Competitor benchmarking: tone, perception created, what to adopt or differentiate.
3) Style guide table: | Element | Recommendation | Example (Before → After) | for Voice, Vocabulary, Sentence length & structure, Emotional resonance, CTA integration.
EVIDENCE (measured metrics and real sample paragraphs):
{{EVIDENCE}}""", "tone"),
    S("semantic_seo", "Semantic SEO map", "Content", "content",
      "Intent & satisfaction goal, PASF, query templates, entity graph, clusters, attributes, context layers and perspectives for a primary query.",
      ["client_page", "competitor_pages", "gsc_queries"],
      """I want you to perform full Semantic SEO analysis for the query: {{primary_query}}
STEP 0 intent (primary intent, user intent statement, search-engine satisfaction goal); STEP 1 50–100 PASF queries; STEP 2 30–50 query templates; STEP 3 entity graph (10–20 primary entities, 5–10 sub-entities each);
STEP 4 pillars, subtopics and semantic cluster groups; STEP 5 50–100 attributes; STEP 6 context layers; STEP 7 15–30 perspectives.
Output in this order: Intent, PASF, Query Templates, Entities, Sub-Entities, Subtopics, Topic Cluster Model, Attributes, Context Layers, Perspectives.
Use the real queries and ranking-page entities in the evidence first.
EVIDENCE:
{{EVIDENCE}}""", "semantic", inputs=["primary_query"]),
    S("geo_query_dataset", "GEO / AI-search query dataset", "AI search", "geo",
      "100+ first-person, conversational AI-search queries across 8 pain models and 9 query modes, in the strict 10-column format.",
      ["gsc_queries"],
      """You are a Senior SEO Strategist and GEO expert. Generate a dataset of 100+ realistic, high-intent queries a real person would ask an AI assistant.
Service: {{primary_keywords}}. Business: {{brand}}. Website: {{url}}. Area: {{location}}. Budget context: {{budget_note}}.
Every query is first-person/conversational and starts with I/My/Is/Can/Do/Should/Need/Find/Searching for/Compare/What/How/Which/Why; 40–50% are 2–3 lines combining budget, area, requirements and safety anxiety.
Cover ALL pain models (trust deficit, uncertainty & trade-offs, dependent & family safety, values & ethics, post-service anxiety, emergency & panic, prevention, regulation & compliance)
and ALL modes (problem-first, symptom-based, AI overview, risk, comparison, trust, transactional-cautious, educational, post-service). No branding, no URLs, no duplicates.
Output ONLY one markdown table with exactly these columns: {{COLUMNS}}
EVIDENCE (real conversational queries already seen in Search Console):
{{EVIDENCE}}""", "geo_dataset", columns=GEO_COLUMNS, inputs=["count"]),
    S("eeat_titles", "E-E-A-T title generator", "Content", "content",
      "White paper & case study titles that prove Experience, Expertise, Authority and Trust — 17-column sheet.",
      [],
      """You are a Senior SEO Strategist and Content Architect. Generate {{count}} high-impact white paper and case study titles.
Core topic: {{topic}}. Industry: {{industry}}. Audience: {{audience}}. Pain points: {{pain_points}}. Solution: {{solution}}. Verifiable outcome: {{metric}}. Year: {{year}}. Tone: {{tone}}.
Mix experience, expertise, authoritativeness and trust formulas evenly; main title ≤70 characters; natural language; year/trend cue; real business pain; no clickbait.
OUTPUT ONLY ONE markdown table with exactly these columns: {{COLUMNS}}""", "eeat_titles", columns=EEAT_COLUMNS, inputs=["topic", "metric", "count"]),
    # ---- more skills
    S("faq_schema", "FAQ & FAQPage schema builder", "AEO", "aeo", "Question queries people actually search, what's missing on the page, and FAQPage JSON-LD.",
      ["client_page", "gsc_queries"],
      """Write answer-first FAQ entries (40–60 words each) for {{url}} using the real question queries below; keep claims factual and on-brand ({{tone}}). Return a markdown table | Question | Answer | then the FAQPage JSON-LD.
EVIDENCE:
{{EVIDENCE}}""", "faq"),
    S("local_landing_pages", "Local landing page planner", "Local", "local", "Service × location page plan with titles, H1s, URLs and the local proof each page needs.",
      ["gsc_queries"], """Plan local landing pages for {{brand}}: services {{primary_keywords}} in locations {{locations}}. For each: target keyword, title, H1, outline, unique local proof. Avoid doorway pages.
EVIDENCE:
{{EVIDENCE}}""", "local_pages", inputs=["locations"]),
    S("blog_titles_from_queries", "Blog titles from search queries", "Content", "content", "Turns the project's real Search Console queries and gaps into blog titles with keywords and format.",
      ["gsc_queries"], """Act as a content strategist. From the real search queries, propose blog titles (≤60 chars) that match intent, each with primary/supporting keywords and format. Markdown table.
EVIDENCE:
{{EVIDENCE}}""", "blog_titles"),
]
BY_KEY = {s["key"]: s for s in SKILLS}
ANALYSES = {"meta_gap": a_meta_gap, "icp": a_icp, "domain_terms": a_domain_terms, "competitor_pages": a_competitor_pages, "keywords": a_keywords,
            "content_plan": a_content_plan, "seo_llm_gap": a_seo_llm_gap, "usp_heading": a_usp_heading, "tone": a_tone, "semantic": a_semantic,
            "geo_dataset": a_geo_dataset, "eeat_titles": a_eeat_titles, "faq": a_faq, "local_pages": a_local_pages}


def library(db) -> list[dict]:
    out = [{k: s[k] for k in ("key", "name", "category", "workstream", "description", "evidence", "output_columns", "inputs", "builtin", "template")}
           | {"version": 1} for s in SKILLS]
    for c in db.scalars(select(mm.PromptSkill).where(mm.PromptSkill.enabled)):
        out.append({"key": c.key, "name": c.name, "category": c.category, "workstream": c.workstream, "description": c.description,
                    "evidence": c.evidence, "output_columns": c.output_columns, "inputs": c.inputs, "builtin": False, "template": c.template,
                    "version": c.version, "id": str(c.id)})
    return out


def get_skill(db, key) -> dict | None:
    if key in BY_KEY:
        return {**BY_KEY[key], "version": 1}
    c = db.scalar(select(mm.PromptSkill).where(mm.PromptSkill.key == key, mm.PromptSkill.enabled))
    if not c:
        return None
    return {"key": c.key, "name": c.name, "category": c.category, "workstream": c.workstream, "description": c.description,
            "evidence": c.evidence or [], "template": c.template, "analysis": None, "output_columns": c.output_columns or [],
            "inputs": c.inputs or [], "builtin": False, "version": c.version}


# ------------------------------------------------------------------ prompt rendering + AI output parsing
def _slim_evidence(ev: dict, budget=24000) -> str:
    def slim_page(p):
        if not isinstance(p, dict) or p.get("status") != 200:
            return p
        q = {k: x for k, x in p.items() if k not in ("text",)}
        q["text_excerpt"] = (p.get("text") or "")[:2500]
        return q
    e = {k: ([slim_page(x) for x in val] if isinstance(val, list) and val and isinstance(val[0], dict) and "url" in val[0] else
             (slim_page(val) if isinstance(val, dict) else val)) for k, val in ev.items()}
    if "gsc_queries" in e:
        e["gsc_queries"] = e["gsc_queries"][:150]
    s = json.dumps(e, default=str, separators=(",", ":"))
    while len(s) > budget and e.get("competitor_pages"):
        for p in e["competitor_pages"]:
            if isinstance(p, dict) and p.get("text_excerpt"):
                p["text_excerpt"] = p["text_excerpt"][: len(p["text_excerpt"]) // 2]
        s = json.dumps(e, default=str, separators=(",", ":"))
        if all(len((p or {}).get("text_excerpt") or "") < 200 for p in e["competitor_pages"]):
            break
    return s[:budget]


def render(template: str, v: dict, ev: dict, columns: list) -> str:
    def val(k):
        x = v.get(k)
        if isinstance(x, list):
            return ", ".join(map(str, x)) or "(not set)"
        return str(x) if x not in (None, "") else "(not set)"
    out = re.sub(r"\{\{\s*([a-zA-Z_]+)\s*\}\}", lambda mt: {"EVIDENCE": _slim_evidence(ev), "COLUMNS": " | ".join(columns)}.get(mt.group(1)) or val(mt.group(1)), template)
    return out


def parse_md_tables(text: str) -> list[list[dict]]:
    tables, cur, header = [], [], None
    for line in (text or "").splitlines() + [""]:
        s = line.strip()
        if s.startswith("|") and s.count("|") >= 2:
            cells = [c.strip() for c in s.strip("|").split("|")]
            if header is None:
                header = cells
            elif all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            else:
                cur.append({header[i] if i < len(header) else f"col{i}": c for i, c in enumerate(cells)})
        else:
            if header and cur:
                tables.append(cur)
            cur, header = [], None
    return tables


# ------------------------------------------------------------------ run
def run_skill(db, project, key: str, inputs: dict | None = None, user_id=None, use_ai: bool = True, run: mm.SkillRun | None = None) -> mm.SkillRun:
    from . import agents
    sk = get_skill(db, key)
    if not sk:
        raise KeyError(key)
    v = merged_inputs(get_profile(db, project), inputs or {})
    if run is None:
        run = mm.SkillRun(organisation_id=project.organisation_id, project_id=project.id, skill_key=key, created_by=user_id)
        db.add(run)
    run.skill_version, run.status = sk["version"], "RUNNING"
    run.inputs = json.loads(json.dumps(v, default=str))
    db.flush()
    ev = collect(db, project, sk["evidence"], v)
    output: dict = {"tables": {}, "notes": []}
    if sk.get("analysis"):
        fn = ANALYSES.get(sk["analysis"])
        res = a_blog_titles(db, project, v, ev) if sk["analysis"] == "blog_titles" else fn(v, ev)
        output["tables"] = {k: t for k, t in res.get("tables", {}).items()}
        output["notes"] += res.get("notes", [])
        if res.get("schema"):
            output["schema"] = res["schema"]
        if res.get("columns"):
            output["columns"] = res["columns"]
    prompt = render(sk["template"], v, ev, sk["output_columns"])
    run.prompt = prompt
    mode = "data"
    if use_ai:
        try:
            text_, meta = agents.llm(db, project.organisation_id, prompt, max_tokens=int((inputs or {}).get("max_tokens") or 4000))
            output["ai_text"] = text_
            ai_tables = parse_md_tables(text_)
            for i, t in enumerate(ai_tables):
                output["tables"][f"AI table {i + 1}" if len(ai_tables) > 1 else "AI result"] = t
            run.model = meta or {}
            mode = "ai+data" if sk.get("analysis") else "ai"
        except agents.NoAI:
            output["notes"].append("AI step skipped: no AI provider connected (Claude, OpenAI, Gemini or Perplexity) — NOT CONNECTED. "
                                   "The data analysis above is complete; the exact prompt is saved so it can run once a provider is connected.")
    has_rows = any(output["tables"].values()) or output.get("ai_text")
    # JSONB does not keep object key order — remember each table's column order explicitly
    output["column_order"] = {name: list(dict.fromkeys(k for r in rows for k in r)) for name, rows in output["tables"].items() if rows}
    run.mode = mode
    run.evidence = json.loads(_slim_evidence(ev, budget=200000) or "{}")
    run.output = json.loads(json.dumps(output, default=str))
    run.status = "COMPLETED" if has_rows else "NEEDS_DATA"
    sizes = ", ".join(f"{k}: {len(t)}" for k, t in output["tables"].items() if t)
    run.summary = (f"{sk['name']} — {sizes or 'no rows'}" + (" · AI narrative included" if output.get("ai_text") else "")
                   + (f" · {output['notes'][0]}" if output["notes"] and not has_rows else ""))[:1000]
    run.finished_at = utcnow()
    return run


# ------------------------------------------------------------------ Excel export
def to_xlsx(run: mm.SkillRun, skill_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for name, rows in (run.output or {}).get("tables", {}).items():
        if not rows:
            continue
        sn = re.sub(r"[\[\]\*\?/\\:]", " ", name)[:31] or "Sheet"
        while sn in used:
            sn = sn[:28] + str(len(used))
        used.add(sn)
        ws = wb.create_sheet(sn)
        cols = ((run.output or {}).get("column_order") or {}).get(name) or list(dict.fromkeys(k for r in rows for k in r.keys()))
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([_cell(r.get(k)) for k in cols])
        for i, k in enumerate(cols, 1):
            width = min(60, max(10, len(str(k)) + 2, max((len(str(r.get(k) or "")) for r in rows[:200]), default=0)))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("About this run")
    meta = [("Skill", skill_name), ("Run", str(run.id)), ("Mode", run.mode), ("Created", str(run.created_at)), ("Summary", run.summary)]
    for n in (run.output or {}).get("notes", []):
        meta.append(("Note", n))
    if (run.output or {}).get("ai_text"):
        meta.append(("AI narrative", run.output["ai_text"][:32000]))
    meta.append(("Prompt used", (run.prompt or "")[:32000]))
    for a, b in meta:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(x):
    if isinstance(x, (list, dict)):
        return json.dumps(x, default=str)[:32000]
    if isinstance(x, str) and x[:1] in "=+-@":
        return "'" + x  # never let user/page text become a spreadsheet formula
    return x


def ordered_output(output: dict) -> dict:
    """Re-apply the stored column order to every table row (for API responses)."""
    out = dict(output or {})
    order = out.get("column_order") or {}
    out["tables"] = {name: [{k: r.get(k) for k in order.get(name, list(r))} for r in rows] for name, rows in (out.get("tables") or {}).items()}
    return out
