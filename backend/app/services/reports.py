"""Audit report export — PDF & Excel, branded / white-label (F-010)."""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select

from .. import models as m
from ..config import get_settings

SEV_COLORS = {"critical": "#B42318", "high": "#D9480F", "medium": "#B08800", "low": "#4B5563"}
GRADE_COLORS = {"pass": "#067647", "fail": "#B42318", "partial": "#B08800", "not_run": "#6B7280", "manual": "#6B7280", "na": "#9CA3AF"}


def _data(db, run_id):
    ar = db.get(m.AuditRun, run_id)
    site = db.get(m.Site, ar.site_id)
    org = db.get(m.Organisation, ar.organisation_id)
    issues = db.scalars(select(m.AuditIssue).where(m.AuditIssue.audit_run_id == run_id).order_by(m.AuditIssue.priority_score.desc())).all()
    results = db.execute(select(m.AuditResult, m.AuditChecklistItem).join(m.AuditChecklistItem, m.AuditResult.checklist_item_id == m.AuditChecklistItem.id)
                         .where(m.AuditResult.audit_run_id == run_id).order_by(m.AuditChecklistItem.code)).all()
    return ar, site, org, issues, results


def _esc(s):
    return (str(s or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(db, run_id, white_label: bool) -> bytes:
    ar, site, org, issues, results = _data(db, run_id)
    b = org.branding or {}
    brand = b.get("company_name") or org.name
    primary = colors.HexColor(b.get("primary_color") or "#1F3A5F")
    buf = io.BytesIO()
    footer = brand if white_label else f"{brand} · Generated with Genranq"

    def on_page(c, d):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#6B7280"))
        c.drawString(18 * mm, 10 * mm, footer)
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {d.page}")
        c.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=f"SEO Audit — {site.domain}", author=brand)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=primary, alignment=0, fontSize=22)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=primary)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#374151"))
    st = []
    st.append(Paragraph(_esc(brand), ParagraphStyle("brand", parent=body, textColor=primary, fontSize=11)))
    st.append(Paragraph(f"SEO Audit Report — {_esc(site.domain)}", h1))
    fin = ar.finished_at.strftime("%d %b %Y %H:%M UTC") if ar.finished_at else "—"
    st.append(Paragraph(f"Completed {fin} · {ar.summary.get('pages_crawled', 0)} pages crawled · "
                        f"check coverage {ar.summary.get('coverage_pct', 0)}%", body))
    st.append(Spacer(1, 6 * mm))
    iss = ar.summary.get("issues", {})
    kpi = [["Overall score", "Critical", "High", "Medium", "Low", "Passed checks"],
           [f"{ar.score if ar.score is not None else '—'}", iss.get("critical", 0), iss.get("high", 0), iss.get("medium", 0),
            iss.get("low", 0), ar.summary.get("pass", 0)]]
    t = Table(kpi, colWidths=[30 * mm] * 6)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica", 8), ("FONT", (0, 1), (-1, 1), "Helvetica-Bold", 16),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6B7280")), ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                           ("LINEBELOW", (0, 1), (-1, 1), 0.5, colors.HexColor("#E5E7EB")), ("BOTTOMPADDING", (0, 1), (-1, 1), 8)]))
    st.append(t)
    st.append(Spacer(1, 6 * mm))
    st.append(Paragraph("Category scores", h2))
    cs = [["Category", "Score"]] + [[c, f"{v}"] for c, v in sorted(ar.category_scores.items(), key=lambda kv: kv[1])]
    t = Table(cs, colWidths=[90 * mm, 30 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB"))]))
    st.append(t)
    st.append(PageBreak())
    st.append(Paragraph("Prioritised issues", h2))
    rows = [["#", "Check", "Severity", "Affected", "Owner", "Priority"]]
    for n, i in enumerate(issues, 1):
        rows.append([n, Paragraph(f"<b>{i.rule}</b> {_esc(i.title)}<br/><font size=7 color='#6B7280'>{_esc(i.evidence.get('message'))}</font>", small),
                     Paragraph(f"<font color='{SEV_COLORS.get(i.severity)}'><b>{i.severity.upper()}</b></font>", small),
                     i.affected_count, i.owner or "", f"{float(i.priority_score):.1f}"])
    t = Table(rows, colWidths=[8 * mm, 95 * mm, 20 * mm, 16 * mm, 18 * mm, 16 * mm], repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB"))]))
    st.append(t)
    st.append(PageBreak())
    st.append(Paragraph("Recommendations", h2))
    for i in issues[:25]:
        st.append(Paragraph(f"<b>{i.rule} — {_esc(i.title)}</b>", body))
        rec = i.evidence.get("ai_recommendation") or i.recommendation
        st.append(Paragraph(_esc(rec).replace("\n", "<br/>"), body))
        if i.affected_urls:
            st.append(Paragraph("Examples: " + ", ".join(_esc(u) for u in i.affected_urls[:3]), small))
        st.append(Spacer(1, 3 * mm))
    st.append(PageBreak())
    st.append(Paragraph("Full checklist", h2))
    rows = [["Code", "Category", "Check", "Source", "Result"]]
    for r, it in results:
        rows.append([it.code, it.category, Paragraph(_esc(it.check), small), Paragraph(_esc(it.data_source), small),
                     Paragraph(f"<font color='{GRADE_COLORS.get(r.grade)}'><b>{r.grade.replace('_', ' ').upper()}</b></font>", small)])
    t = Table(rows, colWidths=[16 * mm, 24 * mm, 80 * mm, 34 * mm, 20 * mm], repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB"))]))
    st.append(t)
    doc.build(st, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def build_xlsx(db, run_id, white_label: bool) -> bytes:
    ar, site, org, issues, results = _data(db, run_id)
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=((org.branding or {}).get("primary_color") or "#1F3A5F").lstrip("#"))

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = hdr, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Summary"
    brand = (org.branding or {}).get("company_name") or org.name
    iss = ar.summary.get("issues", {})
    sheet(ws, ["Metric", "Value"], [
        ["Prepared by", brand if white_label else f"{brand} (Genranq)"], ["Site", site.domain],
        ["Completed", ar.finished_at.strftime("%Y-%m-%d %H:%M UTC") if ar.finished_at else ""],
        ["Overall score", float(ar.score) if ar.score is not None else None], ["Pages crawled", ar.summary.get("pages_crawled")],
        ["Critical issues", iss.get("critical", 0)], ["High issues", iss.get("high", 0)], ["Medium issues", iss.get("medium", 0)],
        ["Low issues", iss.get("low", 0)], ["Passed checks", ar.summary.get("pass", 0)], ["Check coverage %", ar.summary.get("coverage_pct")],
    ] + [[f"Category: {c}", v] for c, v in ar.category_scores.items()], [28, 50])
    sheet(wb.create_sheet("Issues"), ["Priority", "Code", "Category", "Check", "Severity", "Affected", "Owner", "Finding", "Recommendation", "Example URLs"],
          [[float(i.priority_score), i.rule, i.category, i.title, i.severity, i.affected_count, i.owner, i.evidence.get("message"),
            i.evidence.get("ai_recommendation") or i.recommendation, "\n".join(i.affected_urls[:10])] for i in issues],
          [9, 9, 14, 40, 10, 9, 10, 40, 50, 50])
    sheet(wb.create_sheet("Checklist"), ["Code", "Category", "Check", "Data source", "Auto", "Severity", "Fix owner", "Result", "Count", "Finding"],
          [[it.code, it.category, it.check, it.data_source, it.auto, it.severity, it.default_owner, r.grade, r.count, r.message]
           for r, it in results], [9, 14, 45, 22, 6, 9, 11, 10, 8, 60])
    urls = sorted({u for i in issues for u in i.affected_urls})
    sheet(wb.create_sheet("Affected URLs"), ["URL", "Issues"],
          [[u, ", ".join(i.rule for i in issues if u in i.affected_urls)] for u in urls], [80, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def store(db, org_id, filename: str, content: bytes, content_type: str, user_id=None, kind="report") -> m.StoredFile:
    root = get_settings().storage_dir / str(org_id)
    root.mkdir(parents=True, exist_ok=True)
    ref = f"{uuid.uuid4().hex}-{filename}"
    (root / ref).write_bytes(content)
    f = m.StoredFile(organisation_id=org_id, filename=filename, content_type=content_type, size_bytes=len(content),
                     storage_ref=f"{org_id}/{ref}", kind=kind, created_by=user_id)
    db.add(f)
    db.flush()
    return f


def path_of(f: m.StoredFile) -> Path:
    p = (get_settings().storage_dir / f.storage_ref).resolve()
    assert str(p).startswith(str(get_settings().storage_dir.resolve()))
    return p
